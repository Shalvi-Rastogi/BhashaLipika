from docx import Document
from docx.shared import RGBColor, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json
import os
import zipfile
import tempfile
import shutil
import subprocess

from docx.oxml.shared import OxmlElement
from docx.oxml.ns import qn

from PIL import Image, ImageDraw, ImageFont


# ============================================================
# FONT CONFIGURATION
# ============================================================

DEVANAGARI_FONT_PATHS = [
    # macOS
    "/System/Library/Fonts/NotoSansDevanagari-Regular.ttf",
    "/System/Library/Fonts/NotoSansDevanagari-Bold.ttf",
    "/System/Library/Fonts/Supplemental/NotoSansDevanagari-Regular.ttf",
    "/System/Library/Fonts/Supplemental/NotoSansDevanagari-Bold.ttf",

    # Linux
    "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",
    "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Bold.ttf",

    # Windows
    r"C:\Windows\Fonts\Nirmala.ttf",
    r"C:\Windows\Fonts\NirmalaUI.ttf",
    r"C:\Windows\Fonts\Mangal.ttf",
    r"C:\Windows\Fonts\Kokila.ttf",
]

URDU_FONT_PATHS = [
    # macOS
    "/System/Library/Fonts/NotoNaskhArabic-Regular.ttf",
    "/System/Library/Fonts/NotoNaskhArabic-Bold.ttf",
    "/System/Library/Fonts/Supplemental/NotoNaskhArabic-Regular.ttf",
    "/System/Library/Fonts/Supplemental/NotoNaskhArabic-Bold.ttf",

    # Linux
    "/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf",
    "/usr/share/fonts/truetype/noto/NotoNaskhArabic-Bold.ttf",

    # Windows
    r"C:\Windows\Fonts\NotoNaskhArabic-Regular.ttf",
    r"C:\Windows\Fonts\NotoNaskhArabic-Bold.ttf",
    r"C:\Windows\Fonts\Nirmala.ttf",
]


# ============================================================
# SCRIPT DETECTION
# ============================================================

def contains_devanagari(text):
    """
    Check whether text contains Devanagari characters.
    """

    if not text:
        return False

    return any(
        "\u0900" <= char <= "\u097F"
        for char in text
    )


def contains_arabic(text):
    """
    Check whether text contains Arabic/Urdu characters.
    """

    if not text:
        return False

    return any(
        "\u0600" <= char <= "\u06FF"
        for char in text
    )


# ============================================================
# FONT PATH
# ============================================================

def get_font_path(text=""):
    """
    Find a font suitable for the translated script.

    Supports:
    - Hindi / Devanagari
    - Urdu / Arabic
    """

    # Urdu / Arabic
    if contains_arabic(text):

        for font_path in URDU_FONT_PATHS:
            if os.path.exists(font_path):
                return font_path

    # Hindi / Devanagari
    if contains_devanagari(text):

        for font_path in DEVANAGARI_FONT_PATHS:
            if os.path.exists(font_path):
                return font_path

    # Fallback
    for font_path in DEVANAGARI_FONT_PATHS:
        if os.path.exists(font_path):
            return font_path

    for font_path in URDU_FONT_PATHS:
        if os.path.exists(font_path):
            return font_path

    # Try fontconfig if available
    try:

        if shutil.which("fc-match"):

            if contains_arabic(text):
                font_name = "Noto Naskh Arabic"

            elif contains_devanagari(text):
                font_name = "Noto Sans Devanagari"

            else:
                font_name = "Noto Sans"

            result = subprocess.run(
                ["fc-match", "-f", "%{file}", font_name],
                capture_output=True,
                text=True
            )

            path = result.stdout.strip()

            if path and os.path.exists(path):
                return path

    except Exception:
        pass

    return None


def get_image_font_path(text=""):
    """
    Get font for image reconstruction.
    """

    return get_font_path(text)


# ============================================================
# WORD FONT SUPPORT
# ============================================================

def apply_devanagari_font(run, text):
    """
    Set a compatible complex-script font.
    """

    if not text:
        return

    if not (
        contains_devanagari(text)
        or contains_arabic(text)
    ):
        return

    font_name = (
        "Nirmala UI"
        if contains_devanagari(text)
        else "Noto Naskh Arabic"
    )

    rPr = run._element.get_or_add_rPr()

    rFonts = rPr.find(
        qn("w:rFonts")
    )

    if rFonts is None:

        rFonts = OxmlElement("w:rFonts")

        rPr.insert(
            0,
            rFonts
        )

    rFonts.set(
        qn("w:cs"),
        font_name
    )

    rFonts.set(
        qn("w:eastAsia"),
        font_name
    )


# ============================================================
# APPLY RUN STYLE
# ============================================================

def apply_style_to_run(run, style_dict):

    if not style_dict:
        return

    font_name = style_dict.get("font_name")

    is_bold = (
        style_dict.get("bold") is True
        or style_dict.get("cs_bold") is True
    )

    is_italic = (
        style_dict.get("italic") is True
        or style_dict.get("cs_italic") is True
    )

    is_underline = (
        style_dict.get("underline") is True
    )

    # Explicitly apply formatting
    # This prevents formatting from being inherited accidentally.
    run.bold = is_bold
    run.italic = is_italic
    run.underline = is_underline

    # --------------------------------------------------------
    # Font name
    # --------------------------------------------------------

    if font_name:

        try:
            run.font.name = font_name
        except Exception:
            pass

    # --------------------------------------------------------
    # Font size
    # --------------------------------------------------------

    font_size = style_dict.get(
        "font_size"
    )

    if font_size:

        try:
            run.font.size = Pt(
                float(font_size)
            )
        except Exception:
            pass

    # --------------------------------------------------------
    # Font color
    # --------------------------------------------------------

    font_color = style_dict.get(
        "font_color"
    )

    if isinstance(font_color, dict):

        if font_color.get("type") == "rgb":

            hex_value = font_color.get(
                "value"
            )

            if hex_value:

                try:
                    run.font.color.rgb = (
                        RGBColor.from_string(
                            hex_value
                        )
                    )
                except Exception:
                    pass

    # --------------------------------------------------------
    # Complex script formatting
    # --------------------------------------------------------

    rPr = run._element.get_or_add_rPr()

    if is_bold:

        bCs = rPr.find(
            qn("w:bCs")
        )

        if bCs is None:

            bCs = OxmlElement("w:bCs")

            b = rPr.find(
                qn("w:b")
            )

            if b is not None:
                b.addnext(bCs)
            else:
                rPr.insert(
                    0,
                    bCs
                )

    if is_italic:

        iCs = rPr.find(
            qn("w:iCs")
        )

        if iCs is None:

            iCs = OxmlElement("w:iCs")

            i = rPr.find(
                qn("w:i")
            )

            if i is not None:
                i.addnext(iCs)
            else:
                rPr.insert(
                    0,
                    iCs
                )

    if font_size:

        szCs = rPr.find(
            qn("w:szCs")
        )

        if szCs is None:

            szCs = OxmlElement("w:szCs")

            sz = rPr.find(
                qn("w:sz")
            )

            if sz is not None:
                sz.addnext(szCs)
            else:
                rPr.append(szCs)

        try:

            szCs.set(
                qn("w:val"),
                str(
                    int(
                        float(font_size) * 2
                    )
                )
            )

        except Exception:
            pass

    # --------------------------------------------------------
    # Script font
    # --------------------------------------------------------

    apply_devanagari_font(
        run,
        run.text
    )


# ============================================================
# APPLY PARAGRAPH STYLE
# ============================================================

def apply_paragraph_style(
    paragraph,
    paragraph_style
):

    if not paragraph_style:
        return

    alignment = paragraph_style.get(
        "alignment"
    )

    if alignment:

        alignment_map = {

            "left":
                WD_ALIGN_PARAGRAPH.LEFT,

            "center":
                WD_ALIGN_PARAGRAPH.CENTER,

            "right":
                WD_ALIGN_PARAGRAPH.RIGHT,

            "justify":
                WD_ALIGN_PARAGRAPH.JUSTIFY,

        }

        if alignment in alignment_map:

            paragraph.alignment = (
                alignment_map[alignment]
            )


# ============================================================
# ADD TEXT TO PARAGRAPH
# ============================================================

def populate_paragraph(
    paragraph,
    paragraph_data
):

    paragraph.clear()

    paragraph_style = paragraph_data.get(
        "paragraph_style",
        {}
    )

    apply_paragraph_style(
        paragraph,
        paragraph_style
    )

    runs_data = paragraph_data.get(
        "runs",
        []
    )

    if runs_data:

        for run_data in runs_data:

            text = run_data.get(
                "text",
                ""
            )

            if not text:
                continue

            style = run_data.get(
                "style",
                {}
            )

            lines = text.split(
                "\n"
            )

            for index, line in enumerate(lines):

                if line:

                    new_run = paragraph.add_run(
                        line
                    )

                    apply_style_to_run(
                        new_run,
                        style
                    )

                if index < len(lines) - 1:

                    if paragraph.runs:

                        paragraph.runs[-1].add_break()

    else:

        text = paragraph_data.get(
            "text",
            ""
        )

        if text:

            run = paragraph.add_run(
                text
            )

            apply_style_to_run(
                run,
                {}
            )


# ============================================================
# RECONSTRUCT NORMAL PARAGRAPHS
# ============================================================

# def preserve_textbox_runs(paragraph):
#     """
#     Preserve runs that contain Word textboxes/shapes.
#     These runs must not be removed when rebuilding the paragraph.
#     """
#     preserved = []

#     for child in paragraph._p:
#         if child.tag == qn("w:r") and child.xpath(".//w:txbxContent"):
#             preserved.append(deepcopy(child))

#     return preserved

# ============================================================
# TEXTBOX HELPERS
# ============================================================

def clear_paragraph_preserve_textboxes(paragraph):
    """
    Remove normal paragraph runs while preserving
    embedded Word textbox/shape XML.
    """

    for child in list(paragraph._p):

        # Keep paragraph properties
        if child.tag == qn("w:pPr"):
            continue

        # Keep runs that contain textboxes/shapes
        if child.tag == qn("w:r"):
            if any(
                element.tag == qn("w:txbxContent")
                for element in child.iter()
            ):
                continue

            paragraph._p.remove(child)


def reconstruct_textboxes(
    document,
    elements
):
    """
    Replace text inside existing Word textboxes while
    preserving the original shape, border and position.
    """

    textbox_elements = [
        element
        for element in elements
        if element.get("type") == "textbox"
    ]

    if not textbox_elements:
        print("No textbox elements found.")
        return

    translated_textboxes = {
        element.get("textbox_id"): element
        for element in textbox_elements
    }

    original_textboxes = list(
        document._element.iter(
            qn("w:txbxContent")
        )
    )

    print(
        "Original textboxes:",
        len(original_textboxes)
    )

    print(
        "Translated textboxes:",
        len(translated_textboxes)
    )

    for textbox_id, textbox_element in enumerate(
        original_textboxes,
        start=1
    ):

        translated = translated_textboxes.get(
            textbox_id
        )

        if not translated:
            continue

        translated_paragraphs = translated.get(
            "paragraphs",
            []
        )

        textbox_paragraphs = [
            child
            for child in textbox_element
            if child.tag == qn("w:p")
        ]

        for index, textbox_paragraph in enumerate(
            textbox_paragraphs
        ):

            if index >= len(translated_paragraphs):
                continue

            translated_text = str(
                translated_paragraphs[index]
            )

            text_nodes = [
                node
                for node in textbox_paragraph.iter()
                if node.tag == qn("w:t")
            ]

            if not text_nodes:
                continue

            text_nodes[0].text = translated_text

            for node in text_nodes[1:]:
                node.text = ""


# ============================================================
# RECONSTRUCT PARAGRAPHS
# ============================================================

def reconstruct_paragraphs(
    document,
    elements
):

    print(
        "\n========== PARAGRAPH RECONSTRUCTION =========="
    )

    text_elements = [
        element
        for element in elements
        if element.get("type") in [
            "paragraph",
            "heading",
            "list_item"
        ]
    ]

    original_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    print(
        "Original paragraphs:",
        len(original_paragraphs)
    )

    print(
        "Translated text elements:",
        len(text_elements)
    )

    for index, (
        paragraph,
        element
    ) in enumerate(
        zip(
            original_paragraphs,
            text_elements
        )
    ):

        # ----------------------------------------------------
        # Footnote / endnote backup
        # ----------------------------------------------------

        special_tags = []

        for run in paragraph.runs:

            for child in run._element:

                if (
                    child.tag.endswith(
                        "footnoteReference"
                    )
                    or
                    child.tag.endswith(
                        "endnoteReference"
                    )
                ):

                    special_tags.append(
                        child
                    )

        current_text = element.get(
            "text",
            ""
        ).strip()

        # ----------------------------------------------------
        # Existing special Secretary-General handling
        # ----------------------------------------------------

        if (
            current_text.lower()
            in {
                "secretary-general",
                "महासचिव"
            }
            and index > 0
            and text_elements[
                index - 1
            ].get(
                "text",
                ""
            ).strip().endswith(",")
        ):

            continue

        runs_data = element.get(
            "runs",
            []
        )

        paragraph_style = element.get(
            "paragraph_style",
            {}
        )

        clear_paragraph_preserve_textboxes(
            paragraph
        )

        has_textbox = any(
            child.tag == qn("w:r")
            and any(
                node.tag == qn("w:txbxContent")
                for node in child.iter()
            )
            for child in paragraph._p
        )

        if not has_textbox:
            apply_paragraph_style(
                paragraph,
                paragraph_style
            )

        # ----------------------------------------------------
        # Special Secretary-General formatting
        # ----------------------------------------------------

        if (
            current_text.lower()
            in {
                "secretary-general",
                "महासचिव"
            }
            and index > 0
            and text_elements[
                index - 1
            ].get(
                "text",
                ""
            ).strip().endswith(",")
        ):

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.RIGHT
            )

        # ----------------------------------------------------
        # Merge special split paragraph
        # ----------------------------------------------------

        if index + 1 < len(text_elements):

            next_text = text_elements[
                index + 1
            ].get(
                "text",
                ""
            ).strip().lower()

            if (
                current_text.endswith(",")
                and next_text in {
                    "secretary-general",
                    "महासचिव"
                }
            ):

                previous_element = (
                    paragraph._p.getprevious()
                )

                while previous_element is not None:

                    has_text = any(
                        node.tag == qn("w:t")
                        and node.text
                        and node.text.strip()
                        for node in previous_element.iter()
                    )

                    if (
                        has_text
                        or
                        previous_element.tag
                        != qn("w:p")
                    ):
                        break

                    parent = (
                        previous_element.getparent()
                    )

                    parent.remove(
                        previous_element
                    )

                    previous_element = (
                        paragraph._p.getprevious()
                    )

                next_element = text_elements[
                    index + 1
                ]

                runs_data = (
                    runs_data
                    + [
                        {
                            "text": "\n",
                            "style": {}
                        }
                    ]
                    + next_element.get(
                        "runs",
                        []
                    )
                )

                paragraph.paragraph_format.keep_with_next = False
                paragraph.paragraph_format.keep_together = False
                paragraph.paragraph_format.page_break_before = False

                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.RIGHT
                )

                paragraph_properties = (
                    paragraph._p.pPr
                )

                if paragraph_properties is not None:

                    numbering_properties = (
                        paragraph_properties.numPr
                    )

                    if numbering_properties is not None:

                        paragraph_properties.remove(
                            numbering_properties
                        )

                paragraph.paragraph_format.left_indent = None
                paragraph.paragraph_format.right_indent = None
                paragraph.paragraph_format.first_line_indent = None

                if index + 1 < len(original_paragraphs):

                    next_paragraph = (
                        original_paragraphs[
                            index + 1
                        ]
                    )

                    next_paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.RIGHT
                    )

                    next_paragraph.paragraph_format.left_indent = None
                    next_paragraph.paragraph_format.right_indent = None
                    next_paragraph.paragraph_format.first_line_indent = None

        # ----------------------------------------------------
        # Add translated runs
        # ----------------------------------------------------

        if runs_data:

            for run_data in runs_data:

                text = run_data.get(
                    "text",
                    ""
                )

                text = text.replace(
                    ". . f..",
                    "w.e.f."
                )

                if not text:
                    continue

                style = run_data.get(
                    "style",
                    {}
                )

                lines = text.split(
                    "\n"
                )

                for line_index, line in enumerate(lines):

                    if line:

                        new_run = paragraph.add_run(
                            line
                        )

                        apply_style_to_run(
                            new_run,
                            style
                        )

                    if (
                        line_index
                        <
                        len(lines) - 1
                    ):

                        if paragraph.runs:

                            paragraph.runs[-1].add_break()

        else:

            text = element.get(
                "text",
                ""
            )

            if text:

                new_run = paragraph.add_run(
                    text
                )

                apply_style_to_run(
                    new_run,
                    {}
                )

        # ----------------------------------------------------
        # Restore footnote / endnote references
        # ----------------------------------------------------

        if special_tags:

            if not paragraph.runs:

                paragraph.add_run()

            last_run = paragraph.runs[-1]

            for tag in special_tags:

                last_run._element.append(
                    tag
                )


# ============================================================
# GET PARAGRAPH DATA
# ============================================================

def get_paragraph_data_list(
    cell_data
):

    paragraphs = cell_data.get(
        "paragraphs",
        []
    )

    if not isinstance(
        paragraphs,
        list
    ):
        return []

    return paragraphs


# ============================================================
# RECONSTRUCT CELL PARAGRAPHS
# ============================================================

def reconstruct_cell_paragraphs(
    cell,
    cell_data
):

    nested_tables_data = cell_data.get(
        "nested_tables",
        []
    )

    paragraph_data_list = (
        get_paragraph_data_list(
            cell_data
        )
    )

    if nested_tables_data:

        for paragraph in cell.paragraphs:

            if paragraph.text.strip():

                paragraph.clear()

        return

    if not paragraph_data_list:

        text = cell_data.get(
            "text",
            ""
        )

        if text:

            paragraph = (
                cell.paragraphs[0]
                if cell.paragraphs
                else cell.add_paragraph()
            )

            paragraph.clear()

            run = paragraph.add_run(
                text
            )

            apply_style_to_run(
                run,
                {}
            )

        return

    existing_paragraphs = (
        cell.paragraphs
    )

    for index, paragraph_data in enumerate(
        paragraph_data_list
    ):

        if index < len(
            existing_paragraphs
        ):

            paragraph = (
                existing_paragraphs[index]
            )

        else:

            paragraph = (
                cell.add_paragraph()
            )

        populate_paragraph(
            paragraph,
            paragraph_data
        )

    for index in range(
        len(paragraph_data_list),
        len(existing_paragraphs)
    ):

        existing_paragraphs[
            index
        ].clear()


# ============================================================
# RECONSTRUCT NESTED TABLE
# ============================================================

def reconstruct_existing_nested_table(
    existing_table,
    table_data,
    level=1
):

    print(
        f"      Nested table level {level}: "
        f"{len(existing_table.rows)} rows x "
        f"{len(existing_table.columns)} columns"
    )

    rows_data = table_data.get(
        "rows",
        []
    )

    for row_data in rows_data:

        row_index = row_data.get(
            "row_index"
        )

        if (
            row_index is None
            or row_index >= len(
                existing_table.rows
            )
        ):
            continue

        existing_row = (
            existing_table.rows[row_index]
        )

        cells_data = row_data.get(
            "cells",
            []
        )

        for cell_data in cells_data:

            column_index = cell_data.get(
                "column"
            )

            if column_index is None:

                column_index = cell_data.get(
                    "cell_index"
                )

            if (
                column_index is None
                or column_index >= len(
                    existing_row.cells
                )
            ):
                continue

            existing_cell = (
                existing_row.cells[
                    column_index
                ]
            )

            reconstruct_cell_paragraphs(
                existing_cell,
                cell_data
            )

            nested_tables_data = (
                cell_data.get(
                    "nested_tables",
                    []
                )
            )

            if nested_tables_data:

                existing_nested_tables = (
                    existing_cell.tables
                )

                for nested_index, nested_data in enumerate(
                    nested_tables_data
                ):

                    if (
                        nested_index
                        <
                        len(
                            existing_nested_tables
                        )
                    ):

                        reconstruct_existing_nested_table(
                            existing_nested_tables[
                                nested_index
                            ],
                            nested_data,
                            level + 1
                        )


# ============================================================
# RECONSTRUCT TOP-LEVEL TABLE
# ============================================================

def reconstruct_existing_table(
    existing_table,
    table_data
):

    rows_data = table_data.get(
        "rows",
        []
    )

    for row_data in rows_data:

        row_index = row_data.get(
            "row_index"
        )

        if (
            row_index is None
            or row_index >= len(
                existing_table.rows
            )
        ):
            continue

        existing_row = (
            existing_table.rows[row_index]
        )

        cells_data = row_data.get(
            "cells",
            []
        )

        for cell_data in cells_data:

            column_index = cell_data.get(
                "column"
            )

            if column_index is None:

                column_index = cell_data.get(
                    "cell_index"
                )

            if (
                column_index is None
                or column_index >= len(
                    existing_row.cells
                )
            ):
                continue

            existing_cell = (
                existing_row.cells[
                    column_index
                ]
            )

            reconstruct_cell_paragraphs(
                existing_cell,
                cell_data
            )

            nested_tables_data = (
                cell_data.get(
                    "nested_tables",
                    []
                )
            )

            if nested_tables_data:

                original_nested_tables = (
                    existing_cell.tables
                )

                for nested_index, nested_data in enumerate(
                    nested_tables_data
                ):

                    if (
                        nested_index
                        <
                        len(
                            original_nested_tables
                        )
                    ):

                        reconstruct_existing_nested_table(
                            original_nested_tables[
                                nested_index
                            ],
                            nested_data,
                            level=1
                        )


# ============================================================
# TABLE RECONSTRUCTION
# ============================================================

def reconstruct_tables(
    document,
    elements
):

    print(
        "\n========== TABLE RECONSTRUCTION =========="
    )

    table_elements = [
        element
        for element in elements
        if element.get("type") == "table"
    ]

    original_tables = document.tables

    print(
        "Original top-level tables:",
        len(original_tables)
    )

    print(
        "Translated top-level tables:",
        len(table_elements)
    )

    for table_index, table in enumerate(
        original_tables
    ):

        translated_table = None

        for element in table_elements:

            if (
                element.get(
                    "table_index"
                )
                == table_index
            ):

                translated_table = element
                break

        if translated_table is None:

            if table_index < len(
                table_elements
            ):

                translated_table = (
                    table_elements[
                        table_index
                    ]
                )

        if translated_table is None:

            print(
                f"Skipping original table "
                f"{table_index}"
            )

            continue

        print(
            f"\nReconstructing top-level "
            f"table {table_index}"
        )

        reconstruct_existing_table(
            table,
            translated_table
        )


# ============================================================
# IMAGE FONT SIZE
# ============================================================

def calculate_font_size(
    draw,
    text,
    box_width,
    box_height,
    font_path
):

    if (
        not font_path
        or not text
    ):

        return ImageFont.load_default()

    font_size = max(
        8,
        int(box_height * 0.80)
    )

    while font_size >= 8:

        try:

            font = ImageFont.truetype(
                font_path,
                font_size
            )

            bbox = draw.textbbox(
                (0, 0),
                text,
                font=font
            )

            text_width = (
                bbox[2] - bbox[0]
            )

            text_height = (
                bbox[3] - bbox[1]
            )

            if (
                text_width <= box_width
                and text_height <= box_height
            ):

                return font

            font_size -= 1

        except Exception:
            break

    try:

        return ImageFont.truetype(
            font_path,
            8
        )

    except Exception:

        return ImageFont.load_default()


# ============================================================
# BACKGROUND COLOR
# ============================================================

def get_background_color(
    image,
    x1,
    y1,
    x2,
    y2
):

    try:

        pixels = []

        sample_points = [

            (
                max(0, x1 - 2),
                max(0, y1)
            ),

            (
                max(0, x1 - 2),
                min(
                    image.height - 1,
                    y2
                )
            ),

            (
                min(
                    image.width - 1,
                    x2 + 2
                ),
                max(0, y1)
            ),

            (
                min(
                    image.width - 1,
                    x2 + 2
                ),
                min(
                    image.height - 1,
                    y2
                )
            )
        ]

        for point in sample_points:

            pixel = image.getpixel(
                point
            )

            if (
                isinstance(pixel, tuple)
                and len(pixel) >= 3
            ):

                pixels.append(
                    pixel[:3]
                )

        if pixels:

            r = (
                sum(
                    p[0]
                    for p in pixels
                )
                // len(pixels)
            )

            g = (
                sum(
                    p[1]
                    for p in pixels
                )
                // len(pixels)
            )

            b = (
                sum(
                    p[2]
                    for p in pixels
                )
                // len(pixels)
            )

            return (
                r,
                g,
                b
            )

    except Exception:
        pass

    return "white"


# ============================================================
# IMAGE RECONSTRUCTION
# ============================================================

def reconstruct_image(
    original_image_path,
    image_element,
    output_image_path
):

    try:

        image = Image.open(
            original_image_path
        ).convert("RGB")

        draw = ImageDraw.Draw(
            image
        )

        ocr_details = image_element.get(
            "ocr_details",
            []
        )

        if not ocr_details:

            image.save(
                output_image_path
            )

            return True

        for item in ocr_details:

            # ------------------------------------------------
            # IMPORTANT:
            # Translation should create translated_text.
            #
            # If translated_text does not exist, fall back
            # to text so the pipeline does not crash.
            # ------------------------------------------------

            translated_text = item.get(
                "translated_text",
                ""
            )

            if not translated_text:

                translated_text = item.get(
                    "text",
                    ""
                )

            if not translated_text:
                continue

            bbox = item.get(
                "bbox"
            )

            if not bbox:
                continue

            # ------------------------------------------------
            # Bounding box
            # ------------------------------------------------

            x1 = min(
                int(point[0])
                for point in bbox
            )

            y1 = min(
                int(point[1])
                for point in bbox
            )

            x2 = max(
                int(point[0])
                for point in bbox
            )

            y2 = max(
                int(point[1])
                for point in bbox
            )

            x1 = max(
                0,
                x1
            )

            y1 = max(
                0,
                y1
            )

            x2 = min(
                image.width - 1,
                x2
            )

            y2 = min(
                image.height - 1,
                y2
            )

            box_width = max(
                1,
                x2 - x1
            )

            box_height = max(
                1,
                y2 - y1
            )

            # ------------------------------------------------
            # Find correct font
            # ------------------------------------------------

            font_path = (
                get_image_font_path(
                    translated_text
                )
            )

            if font_path:

                print(
                    f"Using image font: "
                    f"{font_path}"
                )

            else:

                print(
                    "No compatible image font found."
                )

            # ------------------------------------------------
            # Cover original text
            # ------------------------------------------------

            background_color = (
                get_background_color(
                    image,
                    x1,
                    y1,
                    x2,
                    y2
                )
            )

            draw.rectangle(
                [
                    x1,
                    y1,
                    x2,
                    y2
                ],
                fill=background_color
            )

            # ------------------------------------------------
            # Calculate font
            # ------------------------------------------------

            font = calculate_font_size(
                draw,
                translated_text,
                box_width,
                box_height,
                font_path
            )

            # ------------------------------------------------
            # Calculate text position
            # ------------------------------------------------

            text_bbox = draw.textbbox(
                (0, 0),
                translated_text,
                font=font
            )

            text_width = (
                text_bbox[2]
                - text_bbox[0]
            )

            text_height = (
                text_bbox[3]
                - text_bbox[1]
            )

            text_x = (
                x1
                + (
                    box_width
                    - text_width
                ) / 2
                - text_bbox[0]
            )

            text_y = (
                y1
                + (
                    box_height
                    - text_height
                ) / 2
                - text_bbox[1]
            )

            # ------------------------------------------------
            # Draw translated text
            # ------------------------------------------------

            draw.text(
                (
                    text_x,
                    text_y
                ),
                translated_text,
                fill="black",
                font=font
            )

        image.save(
            output_image_path
        )

        print(
            f"Reconstructed image saved: "
            f"{output_image_path}"
        )

        return True

    except Exception as e:

        print(
            f"Image reconstruction failed: {e}"
        )

        return False


# ============================================================
# EXTRACT IMAGES FROM DOCX
# ============================================================

def extract_images_from_docx(
    docx_path,
    temp_folder
):

    extracted_images = []

    os.makedirs(
        temp_folder,
        exist_ok=True
    )

    try:

        with zipfile.ZipFile(
            docx_path,
            "r"
        ) as z:

            for file_name in z.namelist():

                if not file_name.startswith(
                    "word/media/"
                ):
                    continue

                image_name = os.path.basename(
                    file_name
                )

                output_path = os.path.join(
                    temp_folder,
                    image_name
                )

                with open(
                    output_path,
                    "wb"
                ) as f:

                    f.write(
                        z.read(
                            file_name
                        )
                    )

                extracted_images.append({

                    "name":
                        image_name,

                    "path":
                        output_path,

                    "zip_name":
                        file_name
                })

        return extracted_images

    except Exception as e:

        print(
            f"Image extraction failed: {e}"
        )

        return []


# ============================================================
# REPLACE IMAGES INSIDE DOCX
# ============================================================

def replace_images_in_docx(
    original_docx_path,
    output_docx_path,
    reconstructed_images
):

    try:

        reconstructed_map = {

            item["original_name"]:
                item["reconstructed_path"]

            for item in reconstructed_images
        }

        temp_docx = (
            output_docx_path
            + ".tmp"
        )

        with zipfile.ZipFile(
            original_docx_path,
            "r"
        ) as zin:

            with zipfile.ZipFile(
                temp_docx,
                "w",
                zipfile.ZIP_DEFLATED
            ) as zout:

                for item in zin.infolist():

                    data = zin.read(
                        item.filename
                    )

                    if item.filename.startswith(
                        "word/media/"
                    ):

                        image_name = (
                            os.path.basename(
                                item.filename
                            )
                        )

                        reconstructed_path = (
                            reconstructed_map.get(
                                image_name
                            )
                        )

                        if (
                            reconstructed_path
                            and
                            os.path.exists(
                                reconstructed_path
                            )
                        ):

                            with open(
                                reconstructed_path,
                                "rb"
                            ) as image_file:

                                data = (
                                    image_file.read()
                                )

                    zout.writestr(
                        item,
                        data
                    )

        os.replace(
            temp_docx,
            output_docx_path
        )

        return True

    except Exception as e:

        print(
            f"Failed to replace images "
            f"in DOCX: {e}"
        )

        return False


# ============================================================
# IMAGE RECONSTRUCTION PIPELINE
# ============================================================

def reconstruct_images_pipeline(
    original_docx_path,
    elements,
    temp_folder
):

    print(
        "\n========== IMAGE RECONSTRUCTION =========="
    )

    image_elements = [
        element
        for element in elements
        if element.get("type")
        == "image_text"
    ]

    print(
        "Translated image elements:",
        len(image_elements)
    )

    if not image_elements:

        print(
            "No image elements found."
        )

        return []

    extracted_images = (
        extract_images_from_docx(
            original_docx_path,
            temp_folder
        )
    )

    if not extracted_images:

        print(
            "No images found inside DOCX."
        )

        return []

    reconstructed_images = []

    for index, image_element in enumerate(
        image_elements
    ):

        image_name = image_element.get(
            "image_name",
            ""
        )

        original_media = image_element.get(
            "original_media",
            ""
        )

        matching_image = None

        # ----------------------------------------------------
        # 1. Match using original_media
        # ----------------------------------------------------

        if original_media:

            original_name = os.path.basename(
                original_media
            )

            matching_image = next(
                (
                    image
                    for image in extracted_images
                    if image["name"]
                    == original_name
                ),
                None
            )

        # ----------------------------------------------------
        # 2. Match using image_name
        # ----------------------------------------------------

        if matching_image is None and image_name:

            matching_image = next(
                (
                    image
                    for image in extracted_images
                    if image["name"]
                    == image_name
                ),
                None
            )

        # ----------------------------------------------------
        # 3. Handle image_1.png vs image1.png
        # ----------------------------------------------------

        if matching_image is None:

            normalized_name = (
                image_name.replace(
                    "_",
                    ""
                )
            )

            matching_image = next(
                (
                    image
                    for image in extracted_images
                    if image["name"].replace(
                        "_",
                        ""
                    )
                    == normalized_name
                ),
                None
            )

        # ----------------------------------------------------
        # 4. Final index-based fallback
        # ----------------------------------------------------

        if (
            matching_image is None
            and index < len(
                extracted_images
            )
        ):

            matching_image = (
                extracted_images[index]
            )

        if matching_image is None:

            print(
                f"Could not match image: "
                f"{image_name}"
            )

            continue

        original_image_path = (
            matching_image["path"]
        )

        reconstructed_image_path = (
            os.path.join(
                temp_folder,
                f"reconstructed_"
                f"{index + 1}_"
                f"{matching_image['name']}"
            )
        )

        success = reconstruct_image(
            original_image_path,
            image_element,
            reconstructed_image_path
        )

        if success:

            reconstructed_images.append({

                "original_name":
                    matching_image["name"],

                "reconstructed_path":
                    reconstructed_image_path
            })

    return reconstructed_images


# ============================================================
# FOOTNOTE TRANSLATION
# ============================================================

def translate_footnotes_xml(
    document
):

    translations = {

        "This Synopsis is not an authoritative record of the proceedings of the Rajya Sabha.":

        "यह सारांश राज्यसभा की कार्यवाही का आधिकारिक प्रामाणिक रिकॉर्ड नहीं है।"
    }

    try:

        for rel in document.part.rels.values():

            if "footnotes" not in rel.reltype:
                continue

            footnote_part = (
                rel.target_part
            )

            if not hasattr(
                footnote_part,
                "_element"
            ):
                continue

            for p in footnote_part._element.xpath(
                ".//w:p"
            ):

                text_nodes = [
                    t
                    for t in p.xpath(
                        ".//w:t"
                    )
                    if t.text
                ]

                if not text_nodes:
                    continue

                full_text = "".join(
                    t.text
                    for t in text_nodes
                )

                for eng, hin in translations.items():

                    if eng in full_text:

                        for t in text_nodes:
                            t.text = ""

                        text_nodes[
                            0
                        ].text = full_text.replace(
                            eng,
                            hin
                        )

    except Exception as e:

        print(
            f"Footnote translation warning: {e}"
        )


# ============================================================
# MAIN DOCX RECONSTRUCTION
# ============================================================

def reconstruct_docx(
    original_docx_path,
    translated_json_path,
    output_docx_path
):

    print(
        "\n=========================================="
    )

    print(
        "STARTING DOCX RECONSTRUCTION"
    )

    print(
        "=========================================="
    )

    # --------------------------------------------------------
    # Validate files
    # --------------------------------------------------------

    if not os.path.exists(
        original_docx_path
    ):

        raise FileNotFoundError(
            "Original DOCX not found: "
            + original_docx_path
        )

    if not os.path.exists(
        translated_json_path
    ):

        raise FileNotFoundError(
            "Translated JSON not found: "
            + translated_json_path
        )

    # --------------------------------------------------------
    # Load original document
    # --------------------------------------------------------

    document = Document(
        original_docx_path
    )

    # --------------------------------------------------------
    # Load translated JSON
    # --------------------------------------------------------

    with open(
        translated_json_path,
        "r",
        encoding="utf-8"
    ) as file:

        translated_data = json.load(
            file
        )

    elements = translated_data.get(
        "elements",
        []
    )

    print(
        "Total JSON elements:",
        len(elements)
    )

    # --------------------------------------------------------
    # Temporary folder
    # --------------------------------------------------------

    base_dir = os.path.dirname(
        os.path.abspath(
            original_docx_path
        )
    )

    temp_folder = os.path.join(
        base_dir,
        "temp",
        os.path.splitext(
            os.path.basename(
                original_docx_path
            )
        )[0]
    )

    os.makedirs(
        temp_folder,
        exist_ok=True
    )

    # --------------------------------------------------------
    # Paragraph reconstruction
    # --------------------------------------------------------

    reconstruct_textboxes(
        document,
        elements
    )

    reconstruct_paragraphs(
        document,
        elements
    )

    # --------------------------------------------------------
    # Table reconstruction
    # --------------------------------------------------------

    reconstruct_tables(
        document,
        elements
    )

    # --------------------------------------------------------
    # Image reconstruction
    #
    # IMPORTANT:
    # This happens before saving the final document.
    # --------------------------------------------------------

    reconstructed_images = (
        reconstruct_images_pipeline(
            original_docx_path,
            elements,
            temp_folder
        )
    )

    # --------------------------------------------------------
    # Footnotes
    # --------------------------------------------------------

    translate_footnotes_xml(
        document
    )

    # --------------------------------------------------------
    # Create output directory
    # --------------------------------------------------------

    output_dir = os.path.dirname(
        os.path.abspath(
            output_docx_path
        )
    )

    os.makedirs(
        output_dir,
        exist_ok=True
    )

    # --------------------------------------------------------
    # Save reconstructed document
    # --------------------------------------------------------

    document.save(
        output_docx_path
    )

    # --------------------------------------------------------
    # Replace reconstructed images
    # --------------------------------------------------------

    if reconstructed_images:

        print(
            "\n========== REPLACING DOCX IMAGES =========="
        )

        success = replace_images_in_docx(
            output_docx_path,
            output_docx_path,
            reconstructed_images
        )

        if success:

            print(
                "Final DOCX updated with "
                "reconstructed images."
            )

    else:

        print(
            "No reconstructed images "
            "to replace."
        )

    # --------------------------------------------------------
    # Done
    # --------------------------------------------------------

    print(
        "\n=========================================="
    )

    print(
        "DOCX RECONSTRUCTION COMPLETED"
    )

    print(
        "=========================================="
    )

    print(
        f"Output: {output_docx_path}"
    )

    return output_docx_path