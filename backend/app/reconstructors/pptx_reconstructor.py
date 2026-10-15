import json
import os
import copy  
import cv2  # NEW: For image inpainting (background cleaning)
import numpy as np  # NEW: For image array manipulation
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Pt

def _shape_at_path(shapes, shape_path):
    current_shapes = shapes
    shape = None
    try:
        for path_position, index in enumerate(shape_path):
            shape = current_shapes[int(index)]
            if path_position < len(shape_path) - 1:
                current_shapes = shape.shapes
    except (AttributeError, IndexError, TypeError, ValueError):
        return None
    return shape

def _shape_path(shape_index):
    if isinstance(shape_index, int):
        return (shape_index,)
    return tuple(str(part) for part in str(shape_index).split("."))

def _replace_text_preserve_format(target, translated_text):
    if hasattr(target, 'paragraphs'):
        text_frame = target
    elif hasattr(target, 'has_text_frame') and target.has_text_frame:
        text_frame = target.text_frame
    else:
        return
        
    translated_lines = str(translated_text).split('\n')
    original_paragraphs = text_frame.paragraphs

    # 1. Line-by-line update while keeping original runs intact
    for i, para in enumerate(original_paragraphs):
        orig_alignment = para.alignment

        # FIX: Asli (non-empty) text run ka BOLD state capture karo
        is_bold = None
        for run in para.runs:
            if run.text.strip():  # Jo run khali nahi hai, uska style pakdo
                if run.font.bold is not None:
                    is_bold = run.font.bold
                break

        if i < len(translated_lines):
            new_text = translated_lines[i]
        else:
            new_text = ""

        if para.runs:
            para.runs[0].text = new_text
            
            # FIX: Bold state force apply karo (agar original me bold tha)
            if is_bold is not None:
                para.runs[0].font.bold = is_bold
                
            for run in para.runs[1:]:
                run.text = ""
        else:
            para.text = new_text

        # Restore the original alignment
        para.alignment = orig_alignment

    # 2. Add extra lines with Soft Returns (\x0b) so bullets don't duplicate
    if len(translated_lines) > len(original_paragraphs) and original_paragraphs:
        extra_lines = translated_lines[len(original_paragraphs):]
        last_para = original_paragraphs[-1]
        
        if last_para.runs:
            last_para.runs[0].text += '\x0b' + '\x0b'.join(extra_lines)
        else:
            last_para.text += '\x0b' + '\x0b'.join(extra_lines)

            
def _set_cell_text(cell, text):
    text_frame = cell.text_frame
    orig_alignment = text_frame.paragraphs[0].alignment if text_frame.paragraphs else None
    lines = str(text).split("\n")

    if text_frame.paragraphs and text_frame.paragraphs[0].runs:
        first_para = text_frame.paragraphs[0]
        first_run = first_para.runs[0]
        first_run.text = lines[0]
        for run in first_para.runs[1:]:
            run.text = ""
        for p in text_frame.paragraphs[1:]:
            p.text = ""

        # Baaki lines ko proper <a:br/> + run me daalo, formatting copy karke
        rpr = first_run._r.find(qn("a:rPr"))
        for line in lines[1:]:
            first_para.add_line_break()
            br = first_para._p.findall(qn("a:br"))[-1]
            new_run = first_para.add_run()
            new_run.text = line
            if rpr is not None:
                br.insert(0, copy.deepcopy(rpr))
                new_run._r.insert(0, copy.deepcopy(rpr))
    else:
        text_frame.text = str(text)

    if orig_alignment:
        text_frame.paragraphs[0].alignment = orig_alignment

    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

def _get_image_font(font_size):
    font_paths = (
        "C:/Windows/Fonts/Nirmala.ttf",
        "C:/Windows/Fonts/NirmalaUI.ttf",
        "C:/Windows/Fonts/Nirmala.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",
        "/System/Library/Fonts/Supplemental/DevanagariMT.ttc",
    )
    for font_path in font_paths:
        if os.path.exists(font_path):
            try:
                return ImageFont.truetype(font_path, font_size)
            except OSError:
                continue
    return ImageFont.load_default()

def _fit_image_text(draw, text, box, max_font_size):
    x1, y1, x2, y2 = box
    box_width = max(1, x2 - x1)
    box_height = max(1, y2 - y1)
    for font_size in range(max_font_size, 5, -1):
        font = _get_image_font(font_size)
        measured = draw.textbbox((0, 0), text, font=font)
        if measured[2] - measured[0] <= box_width and measured[3] - measured[1] <= box_height:
            return font
    return _get_image_font(6)

def reconstruct_image(image_path, ocr_data, translated_text, output_path):
    try:
        image = Image.open(image_path).convert("RGB")
        translated_lines = [line.strip() for line in translated_text.splitlines()]

        # Purana text hatane ke liye: har box ke aas-paas ke background ka median color se fill
        image_cv = cv2.cvtColor(np.array(image), cv2.COLOR_RGB2BGR)
        inpainted_cv = image_cv.copy()
        img_h, img_w = image_cv.shape[:2]
        PAD = 4    # text ke aas-paas kitna erase karna hai
        RING = 8   # background sample karne ke liye bahar ki ring

        for item in ocr_data:
            bbox = item.get("bbox", [])
            if not bbox:
                continue
            xs = [p[0] for p in bbox]
            ys = [p[1] for p in bbox]

            x1 = int(max(0, min(xs) - PAD))
            y1 = int(max(0, min(ys) - PAD))
            x2 = int(min(img_w, max(xs) + PAD))
            y2 = int(min(img_h, max(ys) + PAD))
            if x2 <= x1 or y2 <= y1:
                continue

            ex1 = max(0, x1 - RING)
            ey1 = max(0, y1 - RING)
            ex2 = min(img_w, x2 + RING)
            ey2 = min(img_h, y2 + RING)

            region = image_cv[ey1:ey2, ex1:ex2]
            ring_mask = np.ones(region.shape[:2], dtype=bool)
            ring_mask[y1 - ey1:y2 - ey1, x1 - ex1:x2 - ex1] = False
            ring_pixels = region[ring_mask]
            if ring_pixels.size == 0:
                continue

            bg_color = np.median(ring_pixels, axis=0)
            inpainted_cv[y1:y2, x1:x2] = bg_color.astype(np.uint8)

        image = Image.fromarray(cv2.cvtColor(inpainted_cv, cv2.COLOR_BGR2RGB))
        draw = ImageDraw.Draw(image)

        for index, item in enumerate(ocr_data):
            text = item.get("text", "")
            bbox = item.get("bbox", [])
            if not text or not bbox:
                continue

            xs = [point[0] for point in bbox]
            ys = [point[1] for point in bbox]
            x1, y1, x2, y2 = min(xs), min(ys), max(xs), max(ys)
            box = (x1, y1, x2, y2)

            new_text = item.get("translated_text", "").strip()
            if not new_text:
                new_text = translated_lines[index] if index < len(translated_lines) else ""
            if not new_text.strip():
                continue

            # Background ke hisaab se text color (light bg -> black, dark bg -> white)
            bg_crop = inpainted_cv[int(y1):int(y2), int(x1):int(x2)]
            if bg_crop.size > 0:
                mean_color = cv2.mean(bg_crop)[:3]
                luminance = 0.299 * mean_color[2] + 0.587 * mean_color[1] + 0.114 * mean_color[0]
                text_color = "black" if luminance > 128 else "white"
            else:
                text_color = "black"

            box_height = max(10, y2 - y1)
            font = _fit_image_text(draw, new_text, box, max(10, int(box_height * 0.8)))
            text_bbox = draw.textbbox((0, 0), new_text, font=font)
            text_width = text_bbox[2] - text_bbox[0]
            text_height = text_bbox[3] - text_bbox[1]
            text_x = x1 + max(0, (x2 - x1 - text_width) // 2)
            text_y = y1 + max(0, (y2 - y1 - text_height) // 2) - text_bbox[1]

            draw.text((text_x, text_y), new_text, fill=text_color, font=font)

        image.save(output_path)
        return output_path
    except Exception as e:
        print(f"Error reconstructing image: {e}")
        return image_path
    
def _replace_picture_image(shape, image_path):
    image_part, _ = shape.part.get_or_add_image_part(image_path)
    relationship_id = shape.part.relate_to(image_part, RT.IMAGE)

    blip = shape._element.blipFill.blip
    blip.set(qn("r:embed"), relationship_id)

def _resolve_image_path(image_path, translated_json_path):
    if os.path.exists(image_path):
        return image_path

    normalized_path = image_path.replace("\\", os.sep).replace("/", os.sep)
    image_name = os.path.basename(normalized_path)
    json_directory = os.path.dirname(os.path.abspath(translated_json_path))
    candidates = (
        os.path.join(json_directory, image_name),
        os.path.join(os.path.dirname(json_directory), normalized_path),
    )
    for candidate in candidates:
        if os.path.exists(candidate):
            return candidate
    return image_path

def _table_cell_bounds(table_shape, row_index, column_index):
    left = table_shape.left + sum(
        table_shape.table.columns[index].width for index in range(column_index)
    )
    top = table_shape.top + sum(
        table_shape.table.rows[index].height for index in range(row_index)
    )
    width = table_shape.table.columns[column_index].width
    height = table_shape.table.rows[row_index].height
    return left, top, left + width, top + height

def _fit_nested_tables(slide, slide_width, slide_height):
    table_shapes = [shape for shape in slide.shapes if shape.has_table]
    if len(table_shapes) < 2:
        return

    slide_right = slide_width
    slide_bottom = slide_height
    margin = Pt(2)

    for inner_table in table_shapes:
        center_x = inner_table.left + inner_table.width // 2
        center_y = inner_table.top + inner_table.height // 2
        containing_cells = []

        for outer_table in table_shapes:
            if outer_table is inner_table:
                continue
            for row_index, _ in enumerate(outer_table.table.rows):
                for column_index, _ in enumerate(outer_table.table.columns):
                    bounds = _table_cell_bounds(outer_table, row_index, column_index)
                    if bounds[0] <= center_x <= bounds[2] and bounds[1] <= center_y <= bounds[3]:
                        containing_cells.append(bounds)

        if not containing_cells:
            continue

        cell_left, cell_top, cell_right, cell_bottom = min(
            containing_cells,
            key=lambda bounds: (bounds[2] - bounds[0]) * (bounds[3] - bounds[1]),
        )
        left = max(cell_left + margin, 0)
        top = max(cell_top + margin, 0)
        right = min(cell_right - margin, slide_right)
        bottom = min(cell_bottom - margin, slide_bottom)
        if right > left and bottom > top:
            inner_table.left = left
            inner_table.top = top
            inner_table.width = right - left
            inner_table.height = bottom - top
            if len(inner_table.table.columns) == 2:
                first_column_width = round(inner_table.width * 0.55)
                inner_table.table.columns[0].width = first_column_width
                inner_table.table.columns[1].width = inner_table.width - first_column_width
                inner_table.width = right - left
            row_height = inner_table.height // len(inner_table.table.rows)
            for row_index, row in enumerate(inner_table.table.rows):
                if row_index == len(inner_table.table.rows) - 1:
                    row.height = inner_table.height - row_height * row_index
                else:
                    row.height = row_height

def reconstruct_pptx(original_pptx_path, translated_json_path, output_pptx_path):
    print(f"Loading original PPTX: {original_pptx_path}")
    prs = Presentation(original_pptx_path)
    
    print(f"Loading translated JSON: {translated_json_path}")
    with open(translated_json_path, 'r', encoding='utf-8') as f:
        translated_data = json.load(f)
        
    for slide_index, slide in enumerate(prs.slides):
        slide_id = f"slide_{slide_index + 1}"
        if slide_id not in translated_data:
            continue
            
        shape_dict = {
            str(item['shape_index']): item
            for item in translated_data[slide_id]
        }
        
        for shape_index, item in shape_dict.items():
            shape = _shape_at_path(slide.shapes, _shape_path(shape_index))
            if shape is None:
                continue

            # --- Replace Text Frame Content ---
            if item['type'] == 'text_frame' and shape.has_text_frame:
                translated_text = item.get('translated_text', '')
                if translated_text:
                    _replace_text_preserve_format(shape.text_frame, translated_text)
                    shape.text_frame.word_wrap = True
                    shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

            # --- Replace Table Content ---
            elif item['type'] == 'table' and shape.has_table:
                trans_table = item.get('translated_data', [])
                if trans_table:
                    for row_idx, row in enumerate(shape.table.rows):
                        for col_idx, cell in enumerate(row.cells):
                            if row_idx < len(trans_table) and col_idx < len(trans_table[row_idx]):
                                _set_cell_text(cell, trans_table[row_idx][col_idx])
                                
            # --- Replace Image With Translated Image ---
            elif item['type'] == 'image':
                image_path = item.get('image_path')
                ocr_data = item.get('ocr_data', [])
                translated_text = item.get('translated_text', '')

                if image_path and ocr_data and translated_text:
                    image_path = _resolve_image_path(image_path, translated_json_path)
                    translated_image_path = os.path.join(
                        os.path.dirname(image_path),
                        f"translated_{os.path.basename(image_path)}"
                    )

                    reconstruct_image(image_path, ocr_data, translated_text, translated_image_path)

                    try:
                        _replace_picture_image(shape, translated_image_path)
                    except Exception as e:
                        print(f"Error replacing PPT image: {e}")                                    

        _fit_nested_tables(slide, prs.slide_width, prs.slide_height)

    prs.save(output_pptx_path)
    print(f"Success: Reconstructed PPTX saved to {output_pptx_path}")
    return output_pptx_path