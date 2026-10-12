import os
import json
import hashlib
import posixpath
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
import xml.etree.ElementTree as ET
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import FileResponse
from starlette.background import BackgroundTask
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from app.database.database import get_db
from app.database.models import Document
from app.reconstructors.xlsx_reconstructor import (
    reconstruct_xlsx,
    serialize_xml_preserving_namespaces,
)
from app.reconstructors.docx_reconstructor import reconstruct_docx
from app.reconstructors.pptx_reconstructor import reconstruct_pptx
from app.reconstructors.PDF.native_pdf_reconstructor import reconstruct_native_pdf
from app.utils.xls_converter import convert_xlsx_to_xls
from app.utils.doc_converter import convert_docx_to_doc, _get_soffice_path
from app.utils.rtf_converter import convert_docx_to_rtf


router = APIRouter()

_DOWNLOAD_FORMATS = {
    ".doc": {
        "pdf": ("PDF (.pdf)", "pdf:writer_pdf_Export"),
    },
    ".docx": {
        "doc": ("Legacy Word (.doc)", "doc:MS Word 97"),
        "pdf": ("PDF (.pdf)", "pdf:writer_pdf_Export"),
        "rtf": ("Rich Text (.rtf)", "rtf"),
        "txt": ("Plain Text (.txt)", "txt:Text"),
    },
    ".rtf": {
        "doc": ("Legacy Word (.doc)", "doc:MS Word 97"),
        "docx": ("Word (.docx)", "docx:Office Open XML Text"),
        "pdf": ("PDF (.pdf)", "pdf:writer_pdf_Export"),
        "rtf": ("Rich Text (.rtf)", "rtf"),
        "txt": ("Plain Text (.txt)", "txt:Text"),
    },
    ".ppt": {
        "pdf": ("PDF (.pdf)", "pdf:impress_pdf_Export"),
    },
    ".pptx": {
        "pdf": ("PDF (.pdf)", "pdf:impress_pdf_Export"),
    },
    ".xls": {
        "pdf": ("PDF (.pdf)", "pdf:calc_pdf_Export"),
        "xlsx": ("Excel (.xlsx)", "xlsx:Calc MS Excel 2007 XML"),
    },
    ".xlsx": {
        "pdf": ("PDF (.pdf)", "pdf:calc_pdf_Export"),
        "xls": ("Legacy Excel (.xls)", "xls:MS Excel 97"),
    },
    ".pdf": {
        "pdf": ("PDF (.pdf)", "pdf"),
    },
}

_SPREADSHEET_NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "package_rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
}
_DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _spreadsheet_qname(namespace, tag):
    return f"{{{namespace}}}{tag}"


def _package_relationship_target(source_file, target):
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(
        posixpath.join(posixpath.dirname(source_file), target)
    )


def _prepare_spreadsheet_preview(source_path, preview_dir):
    """Include non-cell drawings in the print area of a preview-only copy."""
    import openpyxl
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.load_workbook(source_path, data_only=False)
    bounds_by_sheet = {}

    for sheet in workbook.worksheets:
        bounds = [None, None, None, None]

        def include_range(min_row, min_column, max_row, max_column):
            bounds[0] = min_row if bounds[0] is None else min(bounds[0], min_row)
            bounds[1] = min_column if bounds[1] is None else min(bounds[1], min_column)
            bounds[2] = max_row if bounds[2] is None else max(bounds[2], max_row)
            bounds[3] = max_column if bounds[3] is None else max(bounds[3], max_column)

        for cell in sheet._cells.values():
            if cell.value is not None:
                include_range(cell.row, cell.column, cell.row, cell.column)

        for image in sheet._images:
            anchor = image.anchor
            start = getattr(anchor, "_from", None)
            end = getattr(anchor, "to", None)
            if start is None:
                continue
            min_row = start.row + 1
            min_column = start.col + 1
            max_row = max(
                min_row,
                end.row + 1 if end is not None else start.row + 1,
            )
            max_column = max(
                min_column,
                end.col + 1 if end is not None else start.col + 1,
            )
            include_range(min_row, min_column, max_row, max_column)

        if bounds[0] is not None:
            bounds_by_sheet[sheet.title] = (
                f"${get_column_letter(bounds[1])}${bounds[0]}:"
                f"${get_column_letter(bounds[3])}${bounds[2]}"
            )

    workbook.close()

    if not bounds_by_sheet:
        return source_path

    main_ns = _SPREADSHEET_NS["main"]
    rel_ns = _SPREADSHEET_NS["r"]
    package_rel_ns = _SPREADSHEET_NS["package_rel"]
    preview_path = os.path.join(
        preview_dir,
        f"{Path(source_path).stem}_preview.xlsx",
    )

    with zipfile.ZipFile(source_path, "r") as source_zip:
        workbook_path = "xl/workbook.xml"
        workbook_root = ET.fromstring(source_zip.read(workbook_path))
        relationships_root = ET.fromstring(
            source_zip.read("xl/_rels/workbook.xml.rels")
        )
        relationship_map = {
            relationship.attrib["Id"]: relationship.attrib["Target"]
            for relationship in relationships_root.findall(
                _spreadsheet_qname(package_rel_ns, "Relationship")
            )
            if relationship.attrib.get("Id") and relationship.attrib.get("Target")
        }

        worksheet_paths = {}
        sheets = workbook_root.find(_spreadsheet_qname(main_ns, "sheets"))
        if sheets is not None:
            for sheet_index, sheet in enumerate(
                sheets.findall(_spreadsheet_qname(main_ns, "sheet"))
            ):
                relationship_id = sheet.attrib.get(
                    _spreadsheet_qname(rel_ns, "id")
                )
                target = relationship_map.get(relationship_id)
                if target:
                    worksheet_paths[sheet.attrib["name"]] = (
                        _package_relationship_target(workbook_path, target),
                        sheet_index,
                    )

        modified_files = {}
        defined_names = workbook_root.find(
            _spreadsheet_qname(main_ns, "definedNames")
        )
        if defined_names is None:
            defined_names = ET.Element(
                _spreadsheet_qname(main_ns, "definedNames")
            )
            calc_properties = workbook_root.find(
                _spreadsheet_qname(main_ns, "calcPr")
            )
            insert_at = (
                list(workbook_root).index(calc_properties)
                if calc_properties is not None
                else len(workbook_root)
            )
            workbook_root.insert(insert_at, defined_names)

        for sheet_name, cell_range in bounds_by_sheet.items():
            worksheet_path, sheet_index = worksheet_paths[sheet_name]
            worksheet_root = ET.fromstring(source_zip.read(worksheet_path))

            sheet_properties = worksheet_root.find(
                _spreadsheet_qname(main_ns, "sheetPr")
            )
            if sheet_properties is None:
                sheet_properties = ET.Element(
                    _spreadsheet_qname(main_ns, "sheetPr")
                )
                worksheet_root.insert(0, sheet_properties)

            page_setup_properties = sheet_properties.find(
                _spreadsheet_qname(main_ns, "pageSetUpPr")
            )
            if page_setup_properties is None:
                page_setup_properties = ET.SubElement(
                    sheet_properties,
                    _spreadsheet_qname(main_ns, "pageSetUpPr"),
                )
            page_setup_properties.set("fitToPage", "1")

            page_setup = worksheet_root.find(
                _spreadsheet_qname(main_ns, "pageSetup")
            )
            if page_setup is None:
                page_setup = ET.Element(
                    _spreadsheet_qname(main_ns, "pageSetup")
                )
                page_margins = worksheet_root.find(
                    _spreadsheet_qname(main_ns, "pageMargins")
                )
                insert_at = (
                    list(worksheet_root).index(page_margins) + 1
                    if page_margins is not None
                    else len(worksheet_root)
                )
                worksheet_root.insert(insert_at, page_setup)
            page_setup.set("fitToWidth", "1")
            page_setup.set("fitToHeight", "0")
            page_setup.set("orientation", "landscape")
            page_setup.attrib.pop("scale", None)

            for existing_name in list(defined_names):
                if (
                    existing_name.attrib.get("name") == "_xlnm.Print_Area"
                    and existing_name.attrib.get("localSheetId") == str(sheet_index)
                ):
                    defined_names.remove(existing_name)

            print_area = ET.SubElement(
                defined_names,
                _spreadsheet_qname(main_ns, "definedName"),
                {
                    "name": "_xlnm.Print_Area",
                    "localSheetId": str(sheet_index),
                },
            )
            quoted_sheet_name = sheet_name.replace("'", "''")
            print_area.text = f"'{quoted_sheet_name}'!{cell_range}"

            modified_files[worksheet_path] = serialize_xml_preserving_namespaces(
                worksheet_root,
                source_zip.read(worksheet_path),
            )

        modified_files[workbook_path] = serialize_xml_preserving_namespaces(
            workbook_root,
            source_zip.read(workbook_path),
        )

        with zipfile.ZipFile(
            preview_path,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as preview_zip:
            for package_item in source_zip.infolist():
                preview_zip.writestr(
                    package_item,
                    modified_files.get(
                        package_item.filename,
                        source_zip.read(package_item.filename),
                    ),
                )

    return preview_path


def _prepare_presentation_preview(source_path, preview_dir):
    preview_path = os.path.join(
        preview_dir,
        f"{Path(source_path).stem}_preview.pptx",
    )
    property_tags = {
        f"{{{_DRAWING_NS}}}rPr",
        f"{{{_DRAWING_NS}}}defRPr",
        f"{{{_DRAWING_NS}}}endParaRPr",
    }
    complex_font_tag = f"{{{_DRAWING_NS}}}cs"
    text_tag = f"{{{_DRAWING_NS}}}t"
    paragraph_tag = f"{{{_DRAWING_NS}}}p"
    run_tag = f"{{{_DRAWING_NS}}}r"
    run_properties_tag = f"{{{_DRAWING_NS}}}rPr"

    with zipfile.ZipFile(source_path, "r") as source_zip:
        with zipfile.ZipFile(
            preview_path,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as preview_zip:
            for package_item in source_zip.infolist():
                content = source_zip.read(package_item.filename)
                if (
                    package_item.filename.startswith("ppt/slides/slide")
                    and package_item.filename.endswith(".xml")
                ):
                    slide_root = ET.fromstring(content)
                    changed = False
                    for paragraph in slide_root.iter(paragraph_tag):
                        for properties in paragraph.iter():
                            if properties.tag not in property_tags:
                                continue
                            text = "".join(
                                node.text or ""
                                for node in properties.iter(text_tag)
                            )
                            if not any(
                                "\u0900" <= character <= "\u097f"
                                for character in text
                            ):
                                continue
                            complex_font = properties.find(complex_font_tag)
                            if complex_font is None:
                                complex_font = ET.SubElement(
                                    properties,
                                    complex_font_tag,
                                )
                            complex_font.set("typeface", "Nirmala UI")
                            changed = True

                        for run in paragraph.iter(run_tag):
                            text = "".join(
                                node.text or "" for node in run.iter(text_tag)
                            )
                            if not any(
                                "\u0900" <= character <= "\u097f"
                                for character in text
                            ):
                                continue

                            run_properties = run.find(run_properties_tag)
                            if run_properties is None:
                                run_properties = ET.Element(run_properties_tag)
                                run.insert(0, run_properties)
                            complex_font = run_properties.find(complex_font_tag)
                            if complex_font is None:
                                complex_font = ET.SubElement(
                                    run_properties,
                                    complex_font_tag,
                                )
                            complex_font.set("typeface", "Nirmala UI")
                            changed = True

                    if changed:
                        content = serialize_xml_preserving_namespaces(
                            slide_root,
                            content,
                        )

                preview_zip.writestr(package_item, content)

    return preview_path


def _get_cached_document_preview(output_path, filename):
    source_stat = os.stat(output_path)
    source_path = str(Path(output_path).resolve())
    cache_key = hashlib.sha256(
        f"{source_path}:{source_stat.st_mtime_ns}:{source_stat.st_size}:preview-render-v4".encode()
    ).hexdigest()
    cache_dir = os.path.join("outputs", ".preview_cache")
    os.makedirs(cache_dir, exist_ok=True)
    cached_path = os.path.join(cache_dir, f"{cache_key}.pdf")

    if os.path.isfile(cached_path):
        return cached_path

    preview_dir = tempfile.mkdtemp(prefix="document-preview-", dir=cache_dir)
    profile_dir = os.path.join(preview_dir, "lo-profile")
    preview_source_path = output_path
    extension = os.path.splitext(filename)[1].lower()

    try:
        if extension == ".xls":
            subprocess.run(
                [
                    _get_soffice_path(),
                    f"-env:UserInstallation={Path(profile_dir).as_uri()}",
                    "--headless",
                    "--convert-to",
                    "xlsx:Calc MS Excel 2007 XML",
                    "--outdir",
                    preview_dir,
                    output_path,
                ],
                check=True,
                capture_output=True,
                text=True,
                timeout=120,
            )
            preview_source_path = os.path.join(
                preview_dir,
                f"{Path(filename).stem}.xlsx",
            )

            if not os.path.isfile(preview_source_path):
                raise RuntimeError("Could not prepare the spreadsheet for preview.")

        if extension in {".xls", ".xlsx"}:
            preview_source_path = _prepare_spreadsheet_preview(
                preview_source_path,
                preview_dir,
            )
        elif extension == ".pptx":
            preview_source_path = _prepare_presentation_preview(
                preview_source_path,
                preview_dir,
            )

        subprocess.run(
            [
                _get_soffice_path(),
                f"-env:UserInstallation={Path(profile_dir).as_uri()}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                preview_dir,
                preview_source_path,
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=120,
        )

        generated_path = os.path.join(
            preview_dir,
            f"{Path(preview_source_path).stem}.pdf",
        )
        if not os.path.isfile(generated_path):
            raise RuntimeError("Office preview conversion produced no PDF.")

        os.replace(generated_path, cached_path)
        return cached_path
    finally:
        shutil.rmtree(preview_dir, ignore_errors=True)


def warm_document_preview(output_path, filename):
    extension = os.path.splitext(filename)[1].lower()
    if extension == ".pdf":
        return output_path
    if extension not in {".doc", ".docx", ".rtf", ".xls", ".xlsx", ".ppt", ".pptx"}:
        raise ValueError(f"Preview is not supported for {extension or 'this file type'}.")
    return _get_cached_document_preview(output_path, filename)


# -----------------------------
# Request Models
# -----------------------------

class CellEdit(BaseModel):
    cell: str
    translated_text: str


class ImageTextEdit(BaseModel):
    sheet_index: int = Field(ge=0)
    image_index: int = Field(ge=0)
    ocr_index: int = Field(ge=0)
    translated_text: str


class DocxTextEdit(BaseModel):
    path: list[str | int]
    translated_text: str


class SaveReviewRequest(BaseModel):
    document_id: int
    edits: list[CellEdit] = Field(default_factory=list)
    image_edits: list[ImageTextEdit] = Field(default_factory=list)
    docx_edits: list[DocxTextEdit] = Field(default_factory=list)
    pptx_edits: list[DocxTextEdit] = Field(default_factory=list)
    pdf_edits: list[DocxTextEdit] = Field(default_factory=list)


def _path_key(container, part):
    if isinstance(container, list):
        if isinstance(part, bool):
            raise TypeError("Boolean values are not valid list indices.")
        return int(part)
    return part


# -----------------------------
# Save Edited Translations
# -----------------------------

@router.post("/review/save")
def save_review(
    request: SaveReviewRequest,
    db: Session = Depends(get_db)
):

    # Find document in database
    document = (
        db.query(Document)
        .filter(Document.id == request.document_id)
        .first()
    )

    if not document:
        raise HTTPException(
            status_code=404,
            detail="Document not found."
        )

    file_type = document.file_type.lower()
    
    if file_type not in {
        "doc", "docx", "rtf", "xls", "xlsx", "ppt", "pptx", "pdf"
    }:
        raise HTTPException(
            status_code=400,
            detail="Review editing supports DOC, DOCX, RTF, XLS, XLSX, PPT, PPTX, and PDF files."
        )

    if not document.output_filename:
        raise HTTPException(
            status_code=404,
            detail="Reconstructed file not found."
        )

    # Path of the original reconstructed file (safe from overwriting)
    original_output_path = os.path.join("outputs", document.output_filename)

    if not os.path.exists(original_output_path):
        raise HTTPException(
            status_code=404,
            detail="Reconstructed document does not exist."
        )

    # Generate new filename appended with '_reviewed'
    name, ext = os.path.splitext(document.output_filename)
    reviewed_filename = f"{name}_reviewed{ext}"
    reviewed_output_path = os.path.join("outputs", reviewed_filename)

    try:
        # ---------------------------------------------
        # <-- BLOCK FOR DOC / DOCX -->
        # ---------------------------------------------
        if file_type in {"doc", "docx", "rtf"}:
            filename_stem = os.path.splitext(
                os.path.basename(document.original_filename)
            )[0]
            language_pair = (
                f"{document.source_language}_to_"
                f"{document.target_language}"
            )
            translated_json_path = os.path.join(
                "outputs",
                f"{filename_stem}_translated_{language_pair}.json",
            )
            if not os.path.exists(translated_json_path):
                translated_json_path = os.path.join(
                    "outputs",
                    f"{filename_stem}_translated_{document.target_language}.json",
                )
            original_docx_path = os.path.join(
                "uploads",
                f"{filename_stem}.docx",
            )

            if not os.path.exists(translated_json_path):
                raise HTTPException(status_code=404, detail="Translated Word document data was not found.")
            if not os.path.exists(original_docx_path):
                raise HTTPException(status_code=404, detail="Original Word document was not found.")

            with open(translated_json_path, "r", encoding="utf-8") as data_file:
                translated_data = json.load(data_file)

            for edit in request.docx_edits:
                target = translated_data
                try:
                    if not edit.path:
                        raise ValueError("Empty edit path.")
                    for part in edit.path[:-1]:
                        target = target[_path_key(target, part)]
                    final_part = _path_key(target, edit.path[-1])
                    if isinstance(target, dict) and final_part == "translated_text":
                        current_value = target.get(final_part, "")
                    else:
                        current_value = target[final_part]
                except (IndexError, KeyError, TypeError, ValueError):
                    raise HTTPException(status_code=400, detail="Invalid Word document text path.")

                # Allow None values to be replaced by new text string
                if current_value is not None and not isinstance(current_value, str):
                    raise HTTPException(status_code=400, detail="Word document edit must target a text value.")
                target[final_part] = edit.translated_text

                if "runs" in edit.path:
                    runs_index = edit.path.index("runs")
                    paragraph = translated_data
                    for part in edit.path[:runs_index]:
                        paragraph = paragraph[_path_key(paragraph, part)]
                    if isinstance(paragraph, dict) and isinstance(paragraph.get("runs"), list):
                        paragraph["text"] = "".join(
                            run.get("text", "") for run in paragraph["runs"] if isinstance(run, dict)
                        )

            with open(translated_json_path, "w", encoding="utf-8") as data_file:
                json.dump(translated_data, data_file, ensure_ascii=False, indent=4)

            if file_type == "rtf":
                conversion_dir = os.path.join(
                    "outputs",
                    f".review_rtf_{document.id}",
                )
                shutil.rmtree(conversion_dir, ignore_errors=True)
                os.makedirs(conversion_dir, exist_ok=True)
                reviewed_docx_path = os.path.join(
                    conversion_dir,
                    f"{filename_stem}_reviewed.docx",
                )
                try:
                    reconstruct_docx(
                        original_docx_path,
                        translated_json_path,
                        reviewed_docx_path,
                    )
                    converted_path = convert_docx_to_rtf(
                        reviewed_docx_path,
                        conversion_dir,
                    )
                    os.replace(converted_path, reviewed_output_path)
                finally:
                    shutil.rmtree(conversion_dir, ignore_errors=True)
            else:
                reconstruct_docx(
                    original_docx_path,
                    translated_json_path,
                    reviewed_output_path,
                )

            if file_type == "doc":
                conversion_dir = os.path.join("outputs", f".review_doc_{document.id}")
                shutil.rmtree(conversion_dir, ignore_errors=True)
                os.makedirs(conversion_dir, exist_ok=True)
                try:
                    converted_path = convert_docx_to_doc(reviewed_output_path, conversion_dir)
                    os.replace(converted_path, reviewed_output_path)
                finally:
                    shutil.rmtree(conversion_dir, ignore_errors=True)

        # ---------------------------------------------
        # <-- BLOCK FOR PPT / PPTX -->
        # ---------------------------------------------
        elif file_type in {"ppt", "pptx"}:
            filename_stem = os.path.splitext(os.path.basename(document.original_filename))[0]
            language_pair = f"{document.source_language}_to_{document.target_language}"
            translated_json_path = os.path.join("outputs", f"{filename_stem}_translated_{language_pair}.json")
            if not os.path.exists(translated_json_path):
                translated_json_path = os.path.join("outputs", f"{filename_stem}_translated_{document.target_language}.json")
            
            original_pptx_path = os.path.join("uploads", os.path.basename(document.original_filename))

            if not os.path.exists(translated_json_path):
                raise HTTPException(status_code=404, detail="Translated Presentation data was not found.")
            if not os.path.exists(original_pptx_path):
                raise HTTPException(status_code=404, detail="Original Presentation was not found.")

            with open(translated_json_path, "r", encoding="utf-8") as data_file:
                translated_data = json.load(data_file)

            for edit in request.pptx_edits:
                target = translated_data
                try:
                    if not edit.path:
                        raise ValueError("Empty edit path.")
                    for part in edit.path[:-1]:
                        target = target[_path_key(target, part)]
                    final_part = _path_key(target, edit.path[-1])
                    if isinstance(target, dict) and final_part == "translated_text":
                        current_value = target.get(final_part, "")
                    else:
                        current_value = target[final_part]
                except (IndexError, KeyError, TypeError, ValueError):
                    raise HTTPException(status_code=400, detail="Invalid Presentation text path.")

                if current_value is not None and not isinstance(current_value, str):
                    raise HTTPException(status_code=400, detail="Presentation edit must target a text value.")
                target[final_part] = edit.translated_text

                if "runs" in edit.path:
                    runs_index = edit.path.index("runs")
                    paragraph = translated_data
                    for part in edit.path[:runs_index]:
                        paragraph = paragraph[_path_key(paragraph, part)]
                    if isinstance(paragraph, dict) and isinstance(paragraph.get("runs"), list):
                        paragraph["text"] = "".join(
                            run.get("text", "") for run in paragraph["runs"] if isinstance(run, dict)
                        )

            with open(translated_json_path, "w", encoding="utf-8") as data_file:
                json.dump(translated_data, data_file, ensure_ascii=False, indent=4)

            # Reconstruct into the new '_reviewed' file
            reconstruct_pptx(
                original_pptx_path,
                translated_json_path,
                reviewed_output_path,
            )

        # ---------------------------------------------
        # <-- BLOCK FOR PDF -->
        # ---------------------------------------------
        elif file_type == "pdf":
            filename_stem = os.path.splitext(os.path.basename(document.original_filename))[0]
            language_pair = f"{document.source_language}_to_{document.target_language}"
            translated_json_path = os.path.join("outputs", f"{filename_stem}_translated_{language_pair}.json")
            if not os.path.exists(translated_json_path):
                translated_json_path = os.path.join("outputs", f"{filename_stem}_translated_{document.target_language}.json")
            
            # Use exact original filename to avoid extension case mismatch
            original_pdf_path = os.path.join("uploads", os.path.basename(document.original_filename))

            if not os.path.exists(translated_json_path):
                raise HTTPException(status_code=404, detail="Translated PDF data was not found.")
            if not os.path.exists(original_pdf_path):
                raise HTTPException(status_code=404, detail="Original PDF was not found.")

            with open(translated_json_path, "r", encoding="utf-8") as data_file:
                translated_data = json.load(data_file)

            # Safely capture edits regardless of what payload key the frontend used
            edits_to_process = request.pdf_edits + request.docx_edits + request.pptx_edits
            print(f"\n[DEBUG PDF] Received {len(edits_to_process)} text edits from frontend.")

            for edit in edits_to_process:
                target = translated_data
                try:
                    if not edit.path:
                        continue
                    for part in edit.path[:-1]:
                        target = target[_path_key(target, part)]
                    final_part = _path_key(target, edit.path[-1])
                    
                    if isinstance(target, dict) and final_part == "translated_text":
                        current_value = target.get(final_part, "")
                    else:
                        current_value = target[final_part]
                except (IndexError, KeyError, TypeError, ValueError):
                    print(f"[DEBUG PDF] Skipping invalid path: {edit.path}")
                    continue

                if current_value is not None and not isinstance(current_value, str):
                    continue
                
                # Apply the new text
                target[final_part] = edit.translated_text

                # CRITICAL FIX: If we edited the main block text, we MUST delete the old segments!
                # Otherwise, the PDF reconstructor will ignore our new text and render the old segments!
                if isinstance(target, dict) and final_part == "translated_text":
                    if "segments" in target:
                        del target["segments"]
                        print(f"[DEBUG PDF] Deleted old 'segments' for block to force new text render.")

            with open(translated_json_path, "w", encoding="utf-8") as data_file:
                json.dump(translated_data, data_file, ensure_ascii=False, indent=4)

            # Reconstruct into the new '_reviewed' file
            print(f"[DEBUG PDF] Reconstructing into {reviewed_output_path}...")
            reconstruct_native_pdf(
                original_pdf_path,
                translated_json_path,
                reviewed_output_path,
                target_language=document.target_language
            )
            print("[DEBUG PDF] PDF successfully reconstructed and saved!\n")

        # ---------------------------------------------
        # <-- BLOCK FOR XLS / IMAGES -->
        # ---------------------------------------------
        elif request.image_edits or file_type == "xls":
            filename_stem = os.path.splitext(os.path.basename(document.original_filename))[0]
            language_pair = f"{document.source_language}_to_{document.target_language}"
            translated_json_path = os.path.join("outputs", f"{filename_stem}_translated_{language_pair}.json")
            if not os.path.exists(translated_json_path):
                translated_json_path = os.path.join("outputs", f"{filename_stem}_translated_{document.target_language}.json")
            
            if file_type == "xls":
                original_path = os.path.join("uploads", f"xls_to_xlsx_{document.id}", f"{filename_stem}.xlsx")
            else:
                original_path = os.path.join("uploads", os.path.basename(document.original_filename))

            if not os.path.exists(translated_json_path):
                raise HTTPException(status_code=404, detail="Translated workbook data was not found.")
            if not os.path.exists(original_path):
                raise HTTPException(status_code=404, detail="Original workbook was not found.")

            with open(translated_json_path, "r", encoding="utf-8") as data_file:
                workbook_data = json.load(data_file)

            sheets = workbook_data.get("sheets", [])
            for edit in request.edits:
                first_sheet_cells = sheets[0].get("cells", []) if sheets else []
                matching_cell = next(
                    (cell for cell in first_sheet_cells if cell.get("coordinate") == edit.cell),
                    None,
                )
                if matching_cell is not None:
                    matching_cell["translated_text"] = edit.translated_text

            for edit in request.image_edits:
                if edit.sheet_index >= len(sheets):
                    raise HTTPException(status_code=400, detail="Invalid worksheet in image edit.")
                images = sheets[edit.sheet_index].get("images", [])
                if edit.image_index >= len(images):
                    raise HTTPException(status_code=400, detail="Invalid image in review edit.")
                ocr_data = images[edit.image_index].get("ocr_data", [])
                if edit.ocr_index >= len(ocr_data):
                    raise HTTPException(status_code=400, detail="Invalid image text in review edit.")
                ocr_data[edit.ocr_index]["translated_text"] = edit.translated_text
                images[edit.image_index]["translated_text"] = "\n".join(
                    item.get("translated_text", "") for item in ocr_data
                )

            with open(translated_json_path, "w", encoding="utf-8") as data_file:
                json.dump(workbook_data, data_file, ensure_ascii=False, indent=4)

            # Reconstruct into the new '_reviewed' file
            reconstruct_xlsx(
                original_file_path=original_path,
                translated_json_path=translated_json_path,
                output_file_path=reviewed_output_path,
            )

            if file_type == "xls":
                conversion_dir = os.path.join("outputs", f".review_xls_{document.id}")
                shutil.rmtree(conversion_dir, ignore_errors=True)
                os.makedirs(conversion_dir, exist_ok=True)
                try:
                    converted_path = convert_xlsx_to_xls(reviewed_output_path, conversion_dir)
                    os.replace(converted_path, reviewed_output_path)
                finally:
                    shutil.rmtree(conversion_dir, ignore_errors=True)
        else:
            # Check latest file to edit (either reviewed or original)
            latest_path = reviewed_output_path if os.path.exists(reviewed_output_path) else original_output_path
            workbook = load_workbook(latest_path)
            worksheet = workbook[workbook.sheetnames[0]]

            for edit in request.edits:
                worksheet[edit.cell] = edit.translated_text

            # Always save into the reviewed file
            workbook.save(reviewed_output_path)

        return {
            "message": "Changes saved successfully.",
            "document_id": document.id,
            "filename": reviewed_filename
        }

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to save changes: {str(e)}"
        )


# -----------------------------
# Download Final Document
# -----------------------------

def _get_final_document_file(document):
    if not document.output_filename:
        raise HTTPException(
            status_code=404,
            detail="Final document not found."
        )

    name, ext = os.path.splitext(document.output_filename)
    reviewed_filename = f"{name}_reviewed{ext}"
    original_output_path = os.path.join("outputs", document.output_filename)
    reviewed_output_path = os.path.join("outputs", reviewed_filename)

    if os.path.exists(reviewed_output_path):
        return reviewed_output_path, reviewed_filename
    if os.path.exists(original_output_path):
        return original_output_path, document.output_filename

    raise HTTPException(
        status_code=404,
        detail="Final document does not exist."
    )


def _download_format_options(filename):
    extension = os.path.splitext(filename)[1].lower()
    formats = _DOWNLOAD_FORMATS.get(extension)
    if formats is None:
        raise HTTPException(
            status_code=400,
            detail=f"Downloads are not supported for {extension or 'this file type'}."
        )

    return [
        {"extension": file_extension, "label": label}
        for file_extension, (label, _) in formats.items()
    ]


@router.get("/review/download-formats/{document_id}")
def get_download_formats(
    document_id: int,
    db: Session = Depends(get_db),
):
    document = db.query(Document).filter(Document.id == document_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found.")

    _, filename = _get_final_document_file(document)
    formats = _download_format_options(filename)
    current_format = os.path.splitext(filename)[1][1:].lower()
    default_format = next(
        (
            option["extension"]
            for option in formats
            if option["extension"] == current_format
        ),
        formats[0]["extension"],
    )
    return {
        "formats": formats,
        "current_format": current_format,
        "default_format": default_format,
    }


@router.get("/review/download/{document_id}")
def download_final_document(
    document_id: int,
    target_format: Optional[str] = Query(default=None, alias="format"),
    db: Session = Depends(get_db)
):

    document = (
        db.query(Document)
        .filter(Document.id == document_id)
        .first()
    )

    if not document:
        raise HTTPException(
            status_code=404,
            detail="Document not found."
        )

    source_path, source_filename = _get_final_document_file(document)
    target_path = source_path
    target_filename = source_filename
    cleanup_dir = None

    if target_format:
        target_format = target_format.lower().lstrip(".")
        source_extension = os.path.splitext(source_filename)[1].lower()
        formats = _DOWNLOAD_FORMATS.get(source_extension)
        if formats is None or target_format not in formats:
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported download format: .{target_format}"
            )

        if target_format != source_extension[1:]:
            target_filename = (
                f"{os.path.splitext(source_filename)[0]}.{target_format}"
            )
            if target_format == "pdf" and source_extension in {".xls", ".xlsx"}:
                target_path = _get_cached_document_preview(
                    source_path,
                    source_filename,
                )
            else:
                _, convert_filter = formats[target_format]
                cleanup_dir = tempfile.mkdtemp(prefix="translated-export-")
                try:
                    subprocess.run(
                        [
                            _get_soffice_path(),
                            "--headless",
                            "--convert-to",
                            convert_filter,
                            "--outdir",
                            cleanup_dir,
                            source_path,
                        ],
                        check=True,
                        capture_output=True,
                        text=True,
                        timeout=120,
                    )
                except FileNotFoundError as error:
                    shutil.rmtree(cleanup_dir, ignore_errors=True)
                    raise HTTPException(
                        status_code=503,
                        detail=str(error),
                    ) from error
                except subprocess.TimeoutExpired as error:
                    shutil.rmtree(cleanup_dir, ignore_errors=True)
                    raise HTTPException(
                        status_code=504,
                        detail="The document conversion timed out.",
                    ) from error
                except subprocess.CalledProcessError as error:
                    shutil.rmtree(cleanup_dir, ignore_errors=True)
                    detail = (error.stderr or "").strip()
                    raise HTTPException(
                        status_code=500,
                        detail=detail or f"Conversion to .{target_format} failed.",
                    ) from error

                target_path = os.path.join(cleanup_dir, target_filename)
                if not os.path.isfile(target_path):
                    shutil.rmtree(cleanup_dir, ignore_errors=True)
                    raise HTTPException(
                        status_code=500,
                        detail=f"Conversion to .{target_format} did not produce a file.",
                    )

    extension_media_types = {
        ".doc": "application/msword",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".odt": "application/vnd.oasis.opendocument.text",
        ".rtf": "application/rtf",
        ".xls": "application/vnd.ms-excel",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".ods": "application/vnd.oasis.opendocument.spreadsheet",
        ".csv": "text/csv",
        ".ppt": "application/vnd.ms-powerpoint",
        ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ".odp": "application/vnd.oasis.opendocument.presentation",
        ".pdf": "application/pdf",
        ".txt": "text/plain",
    }
    extension = os.path.splitext(target_filename)[1].lower()
    media_type = extension_media_types.get(
        extension,
        "application/octet-stream",
    )

    return FileResponse(
        path=target_path,
        filename=target_filename,
        media_type=media_type,
        background=(
            BackgroundTask(shutil.rmtree, cleanup_dir, ignore_errors=True)
            if cleanup_dir
            else None
        ),
    )


@router.get("/review/preview/{document_id}")
def preview_reconstructed_document(
    document_id: int,
    db: Session = Depends(get_db)
):
    document = (
        db.query(Document)
        .filter(Document.id == document_id)
        .first()
    )

    if not document:
        raise HTTPException(
            status_code=404,
            detail="Document not found."
        )

    if not document.output_filename:
        raise HTTPException(
            status_code=404,
            detail="Reconstructed document not found."
        )

    output_path, filename = _get_final_document_file(document)

    extension = os.path.splitext(filename)[1].lower()

    if extension == ".pdf":
        return FileResponse(
            path=output_path,
            filename=filename,
            media_type="application/pdf",
            content_disposition_type="inline",
        )

    if extension not in {
        ".doc", ".docx", ".rtf", ".xls", ".xlsx", ".ppt", ".pptx"
    }:
        raise HTTPException(
            status_code=400,
            detail="Preview is not supported for this document format."
        )

    try:
        preview_path = _get_cached_document_preview(output_path, filename)

        return FileResponse(
            path=preview_path,
            filename=f"{Path(filename).stem}.pdf",
            media_type="application/pdf",
            content_disposition_type="inline",
        )
    except FileNotFoundError as error:
        raise HTTPException(
            status_code=503,
            detail=str(error)
        ) from error
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired, RuntimeError) as error:
        raise HTTPException(
            status_code=500,
            detail=f"Could not prepare document preview: {error}"
        ) from error

def _load_pdf_translated_data(document):
    """Load the translated JSON of a PDF document (same lookup as review/save)."""
    stem = os.path.splitext(os.path.basename(document.original_filename))[0]
    pair = f"{document.source_language}_to_{document.target_language}"

    path = os.path.join("outputs", f"{stem}_translated_{pair}.json")
    if not os.path.exists(path):
        path = os.path.join(
            "outputs", f"{stem}_translated_{document.target_language}.json"
        )
    if not os.path.exists(path):
        raise HTTPException(status_code=404, detail="Translated PDF data was not found.")

    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def _get_pdf_document(document_id, db):
    document = db.query(Document).filter(Document.id == document_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found.")
    if document.file_type.lower() != "pdf":
        raise HTTPException(status_code=400, detail="Not a PDF document.")
    return document


@router.get("/review/pdf-images/{document_id}")
def list_pdf_images(document_id: int, db: Session = Depends(get_db)):
    """List the images of a PDF with their OCR text and file URLs."""
    document = _get_pdf_document(document_id, db)
    data = _load_pdf_translated_data(document)

    result = []
    for page in data.get("pages", []):
        for image in page.get("images", []):
            image_path = image.get("image_path") or ""
            translated_path = os.path.splitext(image_path)[0] + "_translated.png"

            base = (
                f"/review/pdf-image-file/{document_id}/"
                f"{page.get('page_number')}/{image.get('image_index')}"
            )

            blocks = image.get("text_blocks") or []
            result.append({
                "page_number": page.get("page_number"),
                "image_index": image.get("image_index"),
                "original_url": f"{base}/original" if os.path.isfile(image_path) else None,
                "translated_url": f"{base}/translated" if os.path.isfile(translated_path) else None,
                "texts": [
                    {
                        "original": b.get("text", ""),
                        "translated": b.get("translated_text", ""),
                    }
                    for b in blocks
                ],
            })

    return {"images": result}


@router.get("/review/pdf-image-file/{document_id}/{page_number}/{image_index}/{kind}")
def get_pdf_image_file(
    document_id: int,
    page_number: int,
    image_index: int,
    kind: str,
    db: Session = Depends(get_db),
):
    """Serve the original or translated image (looked up by index, not by path)."""
    if kind not in {"original", "translated"}:
        raise HTTPException(status_code=400, detail="kind must be original or translated.")

    document = _get_pdf_document(document_id, db)
    data = _load_pdf_translated_data(document)

    for page in data.get("pages", []):
        if page.get("page_number") != page_number:
            continue
        for image in page.get("images", []):
            if image.get("image_index") != image_index:
                continue

            path = image.get("image_path") or ""
            if kind == "translated":
                path = os.path.splitext(path)[0] + "_translated.png"

            # safety: only serve files inside the outputs folder
            outputs_dir = os.path.abspath("outputs")
            if not os.path.abspath(path).startswith(outputs_dir) or not os.path.isfile(path):
                raise HTTPException(status_code=404, detail="Image file not found.")

            return FileResponse(path)

    raise HTTPException(status_code=404, detail="Image not found.")