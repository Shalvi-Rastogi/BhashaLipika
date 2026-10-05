import json
import os
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .shape_extractor import extract_text_from_shape
from .table_extractor import extract_table_data
from .image_extractor import extract_image


def _extract_shape(shape, shape_path, slide_index, image_output_dir):
    """Extract a shape and recurse into grouped shapes."""
    shape_index = (
        shape_path[0]
        if len(shape_path) == 1
        else ".".join(str(index) for index in shape_path)
    )

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        extracted = []
        for child_index, child in enumerate(shape.shapes):
            extracted.extend(
                _extract_shape(
                    child,
                    (*shape_path, child_index),
                    slide_index,
                    image_output_dir,
                )
            )
        return extracted

    if shape.has_table:
        table_data = extract_table_data(shape, shape_index)
        if table_data:
            # NAYA: Frontend labels ke liye slide_number inject kar rahe hain
            table_data["slide_number"] = slide_index
            return [table_data]
        return []

    if shape.has_text_frame:
        text_data = extract_text_from_shape(shape, shape_index)
        if text_data:
            # NAYA: Frontend labels ke liye slide_number inject kar rahe hain
            text_data["slide_number"] = slide_index
            return [text_data]
        return []

    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        print(f"[PPTX IMAGE DETECTED] slide={slide_index}, shape={shape_index}")

        image_data = extract_image(
            shape,
            shape_index,
            slide_index,
            image_output_dir,
        )

        if image_data:
            print(f"[PPTX IMAGE RESULT] {image_data}")
            # NAYA: Frontend labels ke liye slide_number inject kar rahe hain
            image_data["slide_number"] = slide_index
            return [image_data]
        return []

    return []

def extract_pptx_to_json(pptx_path, output_json_path, image_output_dir="extracted_images"):
    """Main iterator that traverses slides and routes extraction to appropriate modules."""
    prs = Presentation(pptx_path)
    presentation_data = {}

    if not os.path.exists(image_output_dir):
        os.makedirs(image_output_dir)

    for slide_index, slide in enumerate(prs.slides):
        slide_id = f"slide_{slide_index + 1}"
        presentation_data[slide_id] = []

        for shape_index, shape in enumerate(slide.shapes):
            presentation_data[slide_id].extend(
                _extract_shape(
                    shape,
                    (shape_index,),
                    slide_index + 1,
                    image_output_dir,
                )
            )

    with open(output_json_path, 'w', encoding='utf-8') as f:
        json.dump(presentation_data, f, ensure_ascii=False, indent=4)
    
    print(f"Extraction complete. JSON saved to {output_json_path}")
    return output_json_path