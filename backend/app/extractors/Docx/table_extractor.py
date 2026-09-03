from docx.table import Table

from .style_extractor import extract_run_style
from .utils import get_alignment


def _extract_cell_paragraphs(cell):
    """
    Extract all paragraphs inside a table cell,
    including paragraph formatting and run formatting.
    """
    paragraphs = []

    for paragraph_index, paragraph in enumerate(cell.paragraphs):
        runs = []

        for run_index, run in enumerate(paragraph.runs):
            if not run.text:
                continue

            runs.append(
                {
                    "run_index": run_index,
                    "text": run.text,
                    "style": extract_run_style(run),
                }
            )

        paragraphs.append(
            {
                "paragraph_index": paragraph_index,
                "text": paragraph.text,
                "runs": runs,
                "paragraph_style": {
                    "style_name": (
                        paragraph.style.name
                        if paragraph.style
                        else "Normal"
                    ),
                    "alignment": get_alignment(
                        paragraph.alignment
                    ),
                },
            }
        )

    return paragraphs


# ============================================================
# RECURSIVE NESTED TABLE EXTRACTOR
# ============================================================
def _extract_nested_tables(cell):
    """
    Check if a cell contains nested tables and extract them recursively.
    """
    nested_tables = []
    
    if hasattr(cell, "tables") and cell.tables:
        for t_idx, nested_table in enumerate(cell.tables):
            nested_tables.append(_parse_table_block(nested_table, t_idx))
            
    return nested_tables


def _parse_table_block(table_block, table_index):
    """
    Helper to parse any table block (main or nested).
    """
    rows = []

    for row_index, row in enumerate(table_block.rows):
        row_data = {
            "row_index": row_index,
            "cells": []
        }

        for column_index, cell in enumerate(row.cells):
            cell_data = _extract_cell(cell, row_index, column_index)
            row_data["cells"].append(cell_data)

        rows.append(row_data)

    return {
        "table_index": table_index,
        "rows_count": len(table_block.rows),
        "columns_count": len(table_block.columns),
        "rows": rows,
    }


def _extract_cell(cell, row_index, column_index):
    """
    Extract one table cell along with its paragraphs and any nested tables inside it.
    """
    paragraphs = _extract_cell_paragraphs(cell)
    nested_tables = _extract_nested_tables(cell)

    cell_dict = {
        "cell_index": column_index,
        "row": row_index,
        "column": column_index,
        "text": cell.text,
        "original_text": cell.text,
        "paragraphs": paragraphs,
    }

    if nested_tables:
        cell_dict["nested_tables"] = nested_tables

    return cell_dict


def process_table(block, element_id):
    """
    Convert one DOCX table into structured JSON (supports nested tables inside cells).
    """
    if not isinstance(block, Table):
        return None

    table_data = _parse_table_block(block, 0)

    return {
        "id": element_id,
        "type": "table",
        "level": None,
        "order": element_id,
        "rows_count": table_data["rows_count"],
        "columns_count": table_data["columns_count"],
        "rows": table_data["rows"],
    }


def extract_tables(document):
    """
    Backward-compatible helper for callers that need all tables only.
    """
    tables = []

    for table_id, table in enumerate(document.tables, start=1):
        table_data = process_table(table, table_id)
        if table_data is not None:
            tables.append(table_data)

    return tables