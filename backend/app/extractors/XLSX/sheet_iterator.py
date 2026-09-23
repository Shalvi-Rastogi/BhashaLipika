import json
import os
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.utils.cell import range_boundaries

from .text_extractor import extract_cell_text
from .style_extractor import extract_cell_style
from app.extractors.Docx.ocr_extractor import extract_text_from_image


# ============================================================
# XML NAMESPACES
# ============================================================

NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


# ============================================================
# BASIC HELPERS
# ============================================================

def qname(namespace, tag):
    return f"{{{namespace}}}{tag}"


def normalize_zip_path(path):
    """
    Convert paths such as:
        /xl/drawings/drawing1.xml
        ../drawings/drawing1.xml
        drawings/drawing1.xml

    into a clean ZIP path.
    """
    path = path.replace("\\", "/")

    while path.startswith("/"):
        path = path[1:]

    parts = []

    for part in path.split("/"):
        if part == "..":
            if parts:
                parts.pop()
        elif part == "." or part == "":
            continue
        else:
            parts.append(part)

    return "/".join(parts)


def get_relationship_target(source_file, target):
    """
    Resolve an OOXML relationship target relative to source file.
    """

    source_dir = os.path.dirname(source_file)

    if target.startswith("/"):
        return normalize_zip_path(target)

    return normalize_zip_path(
        os.path.join(source_dir, target)
    )


# ============================================================
# TEXTBOX TEXT EXTRACTION
# ============================================================

def extract_text_from_textbox(textbox_shape):
    """
    Extract text from an Excel DrawingML TextBox.

    Structure is generally:

        xdr:sp
            xdr:txBody
                a:p
                    a:r
                        a:t
    """

    tx_body = textbox_shape.find("xdr:txBody", NS)

    if tx_body is None:
        return "", [], []

    paragraphs = []
    runs = []

    for paragraph in tx_body.findall("a:p", NS):

        paragraph_parts = []

        for element in paragraph.iter():

            # Normal text
            if element.tag == qname(NS["a"], "t"):
                text = element.text or ""

                paragraph_parts.append(text)

                if text:
                    runs.append(text)

            # Explicit line break
            elif element.tag == qname(NS["a"], "br"):
                paragraph_parts.append("\n")

        paragraph_text = "".join(paragraph_parts)

        paragraphs.append(paragraph_text)

    full_text = "\n".join(paragraphs)

    return full_text, paragraphs, runs


# ============================================================
# FIND SHEET -> DRAWING RELATIONSHIPS
# ============================================================

def get_sheet_drawing_map(zip_file):
    """
    Returns:

        {
            "Sheet1": [
                "xl/drawings/drawing1.xml"
            ],
            "Sheet2": [
                "xl/drawings/drawing2.xml"
            ]
        }

    This is important because an XLSX drawing belongs to a
    particular worksheet through worksheet relationships.
    """

    result = {}

    workbook_xml_path = "xl/workbook.xml"
    workbook_rels_path = "xl/_rels/workbook.xml.rels"

    if workbook_xml_path not in zip_file.namelist():
        return result

    if workbook_rels_path not in zip_file.namelist():
        return result

    workbook_root = ET.fromstring(
        zip_file.read(workbook_xml_path)
    )

    workbook_rels_root = ET.fromstring(
        zip_file.read(workbook_rels_path)
    )

    # --------------------------------------------------------
    # workbook relationship map
    # --------------------------------------------------------

    workbook_rel_map = {}

    for rel in workbook_rels_root.findall(
        "rel:Relationship",
        NS
    ):
        rel_id = rel.attrib.get("Id")
        target = rel.attrib.get("Target")

        if rel_id and target:
            workbook_rel_map[rel_id] = get_relationship_target(
                workbook_xml_path,
                target
            )

    # --------------------------------------------------------
    # Process worksheets
    # --------------------------------------------------------

    sheets_element = workbook_root.find(
        "main:sheets",
        NS
    )

    if sheets_element is None:
        return result

    for sheet in sheets_element.findall(
        "main:sheet",
        NS
    ):

        sheet_name = sheet.attrib.get("name")

        rel_id = sheet.attrib.get(
            qname(NS["r"], "id")
        )

        if not sheet_name or not rel_id:
            continue

        worksheet_path = workbook_rel_map.get(rel_id)

        if not worksheet_path:
            continue

        worksheet_rels_path = (
            os.path.dirname(worksheet_path)
            + "/_rels/"
            + os.path.basename(worksheet_path)
            + ".rels"
        )

        if worksheet_rels_path not in zip_file.namelist():
            continue

        worksheet_xml = ET.fromstring(
            zip_file.read(worksheet_path)
        )

        worksheet_rels_root = ET.fromstring(
            zip_file.read(worksheet_rels_path)
        )

        worksheet_rel_map = {}

        for rel in worksheet_rels_root.findall(
            "rel:Relationship",
            NS
        ):
            rel_id = rel.attrib.get("Id")
            target = rel.attrib.get("Target")

            if rel_id and target:
                worksheet_rel_map[rel_id] = get_relationship_target(
                    worksheet_path,
                    target
                )

        drawing_paths = []

        for drawing in worksheet_xml.findall(
            "main:drawing",
            NS
        ):

            drawing_rel_id = drawing.attrib.get(
                qname(NS["r"], "id")
            )

            if not drawing_rel_id:
                continue

            drawing_path = worksheet_rel_map.get(
                drawing_rel_id
            )

            if drawing_path:
                drawing_paths.append(drawing_path)

        result[sheet_name] = drawing_paths

    return result


# ============================================================
# EXTRACT TEXT BOXES
# ============================================================

def extract_textboxes_from_xlsx(file_path):
    """
    Extract all modern Excel floating TextBoxes / Shapes
    containing text.

    Returns a dictionary grouped by worksheet.
    """

    textboxes_by_sheet = {}

    with zipfile.ZipFile(
        file_path,
        "r"
    ) as zip_file:

        sheet_drawing_map = get_sheet_drawing_map(
            zip_file
        )

        for sheet_name, drawing_paths in sheet_drawing_map.items():

            textboxes = []

            for drawing_path in drawing_paths:

                if drawing_path not in zip_file.namelist():
                    continue

                try:
                    drawing_root = ET.fromstring(
                        zip_file.read(drawing_path)
                    )

                except ET.ParseError:
                    print(
                        f"[XLSX WARNING] Could not parse drawing: "
                        f"{drawing_path}"
                    )
                    continue

                # ------------------------------------------------
                # Every xdr:sp can represent a shape.
                # We only select shapes containing txBody.
                # ------------------------------------------------

                for shape in drawing_root.findall(
                    ".//xdr:sp",
                    NS
                ):

                    tx_body = shape.find(
                        "xdr:txBody",
                        NS
                    )

                    if tx_body is None:
                        continue

                    text, paragraphs, runs = (
                        extract_text_from_textbox(shape)
                    )

                    if not text.strip():
                        continue

                    # ------------------------------------------------
                    # Shape ID and name
                    # ------------------------------------------------

                    c_nv_pr = shape.find(
                        "xdr:nvSpPr/xdr:cNvPr",
                        NS
                    )

                    shape_id = ""
                    shape_name = ""

                    if c_nv_pr is not None:
                        shape_id = c_nv_pr.attrib.get(
                            "id",
                            ""
                        )

                        shape_name = c_nv_pr.attrib.get(
                            "name",
                            ""
                        )

                    textbox_info = {
                        "type": "textbox",

                        "sheet_name": sheet_name,

                        "drawing_file": drawing_path,

                        "shape_id": shape_id,

                        "shape_name": shape_name,

                        "original_text": text,

                        "text": text,

                        "paragraphs": paragraphs,

                        "runs": runs
                    }

                    textboxes.append(
                        textbox_info
                    )

                    print(
                        f"[XLSX TEXTBOX] "
                        f"{sheet_name} | "
                        f"{shape_name} | "
                        f"{text[:100]}"
                    )

            if textboxes:
                textboxes_by_sheet[sheet_name] = textboxes

    return textboxes_by_sheet


# ============================================================
# TABLE EXTRACTION
# ============================================================

def extract_sheet_tables(sheet):
    """Extract structured metadata and text for every Excel table."""

    tables = []

    for table in sheet.tables.values():
        min_column, min_row, max_column, max_row = range_boundaries(
            table.ref
        )

        style_info = table.tableStyleInfo
        table_info = {
            "type": "table",
            "name": table.name,
            "display_name": table.displayName,
            "ref": table.ref,
            "header_row_count": table.headerRowCount,
            "totals_row_count": table.totalsRowCount,
            "auto_filter": table.autoFilter is not None,
            "style": {
                "name": style_info.name if style_info else None,
                "show_first_column": (
                    style_info.showFirstColumn if style_info else False
                ),
                "show_last_column": (
                    style_info.showLastColumn if style_info else False
                ),
                "show_row_stripes": (
                    style_info.showRowStripes if style_info else False
                ),
                "show_column_stripes": (
                    style_info.showColumnStripes if style_info else False
                ),
            },
            "columns": [],
            "cells": [],
        }

        for column in table.tableColumns:
            table_info["columns"].append({
                "id": column.id,
                "name": column.name,
            })

        for row in sheet.iter_rows(
            min_row=min_row,
            max_row=max_row,
            min_col=min_column,
            max_col=max_column,
        ):
            for cell in row:
                text = extract_cell_text(cell)
                table_info["cells"].append({
                    "type": "table_cell",
                    "coordinate": cell.coordinate,
                    "row": cell.row,
                    "column": cell.column,
                    "original_text": text,
                    "text": text,
                })

        tables.append(table_info)

    return tables


# ============================================================
# MAIN XLSX EXTRACTION
# ============================================================

def extract_images_from_sheet(sheet, sheet_name, output_dir):
    """Extract images from an Excel sheet and run OCR on them."""

    images_data = []

    image_output_dir = os.path.join(
        output_dir,
        f"{sheet_name}_images"
    )

    os.makedirs(image_output_dir, exist_ok=True)

    for image_index, image in enumerate(sheet._images):

        try:
            image_bytes = image._data()

            image_filename = (
                f"{sheet_name}_image_{image_index + 1}.png"
            )

            image_path = os.path.join(
                image_output_dir,
                image_filename
            )

            # Save image
            with open(image_path, "wb") as f:
                f.write(image_bytes)

            print(f"[XLSX IMAGE] Saved: {image_path}")

            # Run OCR
            ocr_results = extract_text_from_image(image_path)

            print(f"[XLSX OCR] {ocr_results}")

            extracted_text = "\n".join(
                item["text"]
                for item in ocr_results
                if item.get("text")
            )

            # Get image position
            row = None
            column = None

            anchor = image.anchor

            if hasattr(anchor, "_from"):
                row = anchor._from.row + 1
                column = anchor._from.col + 1

            images_data.append({
                "type": "image",
                "image_index": image_index,
                "image_path": image_path,
                "row": row,
                "column": column,
                "original_text": extracted_text,
                "ocr_data": ocr_results,
                "translated_text": ""
            })

        except Exception as e:
            print(
                f"[XLSX IMAGE ERROR] "
                f"Sheet={sheet_name}, "
                f"Image={image_index}: {e}"
            )

    return images_data

def extract_xlsx_to_json(file_path, json_path):

    print("\n========================================")
    print("XLSX EXTRACTION STARTED")
    print("========================================")

    # --------------------------------------------------------
    # Open workbook only for reading.
    # IMPORTANT:
    # We never save this workbook here.
    # --------------------------------------------------------

    wb = openpyxl.load_workbook(
        file_path,
        data_only=False
    )

    document_data = {
        "document_name": os.path.basename(file_path),

        "document_type": "xlsx",

        "sheets": []
    }

    total_cells = 0
    total_textboxes = 0

    # ========================================================
    # 1. NORMAL CELL EXTRACTION
    # ========================================================

    for sheet_name in wb.sheetnames:

        sheet = wb[sheet_name]

        sheet_data = {
            "sheet_name": sheet_name,

            "cells": [],

            "tables": extract_sheet_tables(sheet),

            "textboxes": [],

            "images": []
        }

        for row in sheet.iter_rows():

            for cell in row:

                text = extract_cell_text(cell)

                if text:

                    cell_info = {
                        "type": "cell",

                        "coordinate": cell.coordinate,

                        "row": cell.row,

                        "column": cell.column,

                        "original_text": text,

                        "text": text,

                        "style": extract_cell_style(cell)
                    }

                    sheet_data["cells"].append(
                        cell_info
                    )

                    total_cells += 1

        images = extract_images_from_sheet(
            sheet,
            sheet_name,
            os.path.dirname(json_path)
        )

        sheet_data["images"] = images

        document_data["sheets"].append(
            sheet_data
        )

    wb.close()

    # ========================================================
    # 2. FLOATING TEXTBOX EXTRACTION
    # ========================================================

    textboxes_by_sheet = extract_textboxes_from_xlsx(
        file_path
    )

    for sheet_data in document_data["sheets"]:

        sheet_name = sheet_data["sheet_name"]

        textboxes = textboxes_by_sheet.get(
            sheet_name,
            []
        )

        sheet_data["textboxes"] = textboxes

        total_textboxes += len(textboxes)

    # ========================================================
    # 3. SAVE JSON
    # ========================================================

    with open(
        json_path,
        "w",
        encoding="utf-8"
    ) as f:

        json.dump(
            document_data,
            f,
            indent=4,
            ensure_ascii=False
        )

    print("\n========================================")
    print("XLSX EXTRACTION COMPLETED")
    print("========================================")

    print(
        f"[XLSX] Normal cells extracted : {total_cells}"
    )

    print(
        f"[XLSX] TextBoxes extracted     : {total_textboxes}"
    )

    if total_textboxes == 0:

        print(
            "[XLSX WARNING] No floating TextBox was detected."
        )

        print(
            "[XLSX WARNING] Make sure the Excel object is "
            "Insert -> Text Box / Shape."
        )

    print(
        f"[XLSX] JSON saved to: {json_path}"
    )

    return document_data